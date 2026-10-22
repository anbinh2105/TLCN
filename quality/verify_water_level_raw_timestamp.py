import io
import os
from datetime import datetime

from dotenv import load_dotenv
from minio import Minio
from openpyxl import load_workbook


load_dotenv()

client = Minio(
    os.getenv("MINIO_ENDPOINT", "localhost:9000"),
    access_key=os.getenv("MINIO_ACCESS_KEY", "admin"),
    secret_key=os.getenv("MINIO_SECRET_KEY", "admin12345"),
    secure=os.getenv("MINIO_SECURE", "false").lower() == "true",
)

bucket = os.getenv("MINIO_BUCKET", "flood-lakehouse")
prefix = "bronze/water_level/"

objects = [
    obj.object_name
    for obj in client.list_objects(
        bucket,
        prefix=prefix,
        recursive=True,
    )
    if obj.object_name.lower().endswith(".xlsx")
]

if not objects:
    raise RuntimeError("No Water Level XLSX found.")

object_name = objects[0]

print("=" * 80)
print("VERIFY RAW WATER LEVEL TIMESTAMPS")
print("=" * 80)
print("Object:", object_name)

response = client.get_object(bucket, object_name)

try:
    data = response.read()
finally:
    response.close()
    response.release_conn()

wb = load_workbook(
    io.BytesIO(data),
    data_only=True,
    read_only=True,
)

ws = wb["Data"]

misaligned = []
total = 0

for excel_row, row in enumerate(
    ws.iter_rows(min_row=2, values_only=True),
    start=2,
):
    timestamp = row[4]  # Column E
    value = row[5]      # Column F

    if timestamp is None:
        continue

    total += 1

    if not isinstance(timestamp, datetime):
        print(
            f"[NON-DATETIME] row={excel_row} "
            f"value={repr(timestamp)}"
        )
        continue

    if (
        timestamp.minute % 15 != 0
        or timestamp.second != 0
    ):
        misaligned.append(
            (
                excel_row,
                timestamp,
                value,
            )
        )


print()
print(f"Total RAW timestamps      : {total}")
print(f"RAW misaligned timestamps : {len(misaligned)}")

print()
print("First 50 RAW misaligned timestamps:")
print("-" * 80)

for excel_row, timestamp, value in misaligned[:50]:
    print(
        f"Excel row={excel_row:<8} "
        f"timestamp={timestamp} "
        f"water_level={value}"
    )


# ------------------------------------------------------------
# Explicitly search timestamps found in Silver
# ------------------------------------------------------------

targets = {
    datetime(2024, 5, 26, 2, 25),
    datetime(2024, 5, 26, 17, 10),
    datetime(2024, 5, 27, 3, 25),
    datetime(2024, 5, 28, 4, 5),
    datetime(2024, 5, 28, 18, 25),
}

found = []

for excel_row, row in enumerate(
    ws.iter_rows(min_row=2, values_only=True),
    start=2,
):
    timestamp = row[4]

    if timestamp in targets:
        found.append(
            (
                excel_row,
                timestamp,
                row[5],
            )
        )


print()
print("=" * 80)
print("TARGET TIMESTAMPS FOUND DIRECTLY IN RAW XLSX")
print("=" * 80)

if not found:
    print("NONE")
else:
    for excel_row, timestamp, value in found:
        print(
            f"row={excel_row} "
            f"timestamp={timestamp} "
            f"value={value}"
        )


wb.close()

print()
print("=" * 80)
print("VERIFY COMPLETE")
print("=" * 80)