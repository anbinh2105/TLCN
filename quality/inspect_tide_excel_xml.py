"""
quality/inspect_tide_excel_xml.py

Purpose
-------
Investigate the abnormal Tide height cells in 2021 XLSX files.

This script DOES NOT modify Bronze data.

For each 2021 Tide XLSX object in MinIO Bronze, it compares:

    Excel raw XML <v>
        vs.
    OpenPyXL interpreted value
        vs.
    Excel number format

This is used to determine whether Tide height values that OpenPyXL
interprets as datetime can be deterministically recovered.

Architecture
------------
MinIO Bronze XLSX
        |
        +--> OpenPyXL interpretation
        |
        +--> raw XLSX XML inspection
        |
        +--> comparison report

Output
------
reports/quality/tide_excel_xml_inspection.txt
"""

from __future__ import annotations

import io
import os
import re
import zipfile
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

from dotenv import load_dotenv
from minio import Minio
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


# ============================================================
# PROJECT CONFIG
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

load_dotenv(PROJECT_ROOT / ".env")

MINIO_ENDPOINT = os.getenv(
    "MINIO_ENDPOINT",
    "localhost:9000",
)

MINIO_ACCESS_KEY = os.getenv(
    "MINIO_ACCESS_KEY",
    "admin",
)

MINIO_SECRET_KEY = os.getenv(
    "MINIO_SECRET_KEY",
    "admin12345",
)

MINIO_BUCKET = os.getenv(
    "MINIO_BUCKET",
    "flood-lakehouse",
)

MINIO_SECURE = (
    os.getenv("MINIO_SECURE", "false")
    .strip()
    .lower()
    == "true"
)

BRONZE_PREFIX = "bronze/tide/"

REPORT_DIR = PROJECT_ROOT / "reports" / "quality"

REPORT_PATH = (
    REPORT_DIR
    / "tide_excel_xml_inspection.txt"
)


# ============================================================
# XML NAMESPACES
# ============================================================

MAIN_NS = (
    "http://schemas.openxmlformats.org/"
    "spreadsheetml/2006/main"
)

REL_NS = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/relationships"
)

PACKAGE_REL_NS = (
    "http://schemas.openxmlformats.org/"
    "package/2006/relationships"
)

NS = {
    "main": MAIN_NS,
    "rel": REL_NS,
    "pkg": PACKAGE_REL_NS,
}


# ============================================================
# MINIO CLIENT
# ============================================================

minio_client = Minio(
    MINIO_ENDPOINT,
    access_key=MINIO_ACCESS_KEY,
    secret_key=MINIO_SECRET_KEY,
    secure=MINIO_SECURE,
)


# ============================================================
# REPORT WRITER
# ============================================================

class Reporter:

    def __init__(self):
        self.lines: list[str] = []

    def write(self, text=""):
        text = str(text)

        print(text)

        self.lines.append(text)

    def section(self, title):
        self.write()
        self.write("=" * 100)
        self.write(title)
        self.write("=" * 100)

    def subsection(self, title):
        self.write()
        self.write("-" * 100)
        self.write(title)
        self.write("-" * 100)

    def save(self):
        REPORT_DIR.mkdir(
            parents=True,
            exist_ok=True,
        )

        REPORT_PATH.write_text(
            "\n".join(self.lines),
            encoding="utf-8",
        )


report = Reporter()


# ============================================================
# HELPERS
# ============================================================

def extract_year_from_filename(
    object_name: str,
) -> Optional[int]:

    filename = Path(object_name).name

    match = re.search(
        r"(20\d{2})",
        filename,
    )

    if not match:
        return None

    return int(match.group(1))


def extract_month_from_filename(
    object_name: str,
) -> Optional[int]:

    filename = Path(object_name).name.lower()

    match = re.search(
        r"thang\s*([0-9]{1,2})",
        filename,
    )

    if not match:
        match = re.search(
            r"thang([0-9]{1,2})",
            filename,
        )

    if not match:
        return None

    month = int(match.group(1))

    if 1 <= month <= 12:
        return month

    return None


def classify_value(value):

    if value is None:
        return "NULL"

    if isinstance(value, datetime):
        return "DATETIME"

    if isinstance(value, date):
        return "DATE"

    if isinstance(value, time):
        return "TIME"

    if isinstance(value, bool):
        return "BOOLEAN"

    if isinstance(value, (int, float)):
        return "NUMERIC"

    if isinstance(value, str):

        stripped = value.strip()

        if not stripped:
            return "EMPTY_STRING"

        try:
            float(stripped)
            return "STRING_NUMERIC"
        except ValueError:
            return "STRING"

    return type(value).__name__.upper()


def download_object(
    object_name: str,
) -> bytes:

    response = minio_client.get_object(
        MINIO_BUCKET,
        object_name,
    )

    try:
        return response.read()

    finally:
        response.close()
        response.release_conn()


# ============================================================
# XLSX XML HELPERS
# ============================================================

def normalize_zip_path(path: str) -> str:

    path = path.replace("\\", "/")

    if path.startswith("/"):
        path = path[1:]

    if path.startswith("xl/"):
        return path

    return "xl/" + path


def get_first_worksheet_xml_path(
    xlsx_bytes: bytes,
) -> str:
    """
    Resolve the first worksheet XML path through workbook.xml
    and workbook relationships instead of assuming sheet1.xml.
    """

    with zipfile.ZipFile(
        io.BytesIO(xlsx_bytes)
    ) as zf:

        workbook_xml = zf.read(
            "xl/workbook.xml"
        )

        workbook_root = ET.fromstring(
            workbook_xml
        )

        sheets = workbook_root.find(
            f"{{{MAIN_NS}}}sheets"
        )

        if sheets is None or len(sheets) == 0:
            raise ValueError(
                "Workbook contains no worksheets."
            )

        first_sheet = sheets[0]

        relationship_id = first_sheet.attrib.get(
            f"{{{REL_NS}}}id"
        )

        sheet_name = first_sheet.attrib.get(
            "name"
        )

        if not relationship_id:
            raise ValueError(
                f"Cannot find relationship ID "
                f"for worksheet {sheet_name!r}."
            )

        rels_xml = zf.read(
            "xl/_rels/workbook.xml.rels"
        )

        rels_root = ET.fromstring(
            rels_xml
        )

        target = None

        for relationship in rels_root:

            if (
                relationship.attrib.get("Id")
                == relationship_id
            ):
                target = relationship.attrib.get(
                    "Target"
                )
                break

        if not target:
            raise ValueError(
                f"Cannot resolve XML path for "
                f"worksheet {sheet_name!r}."
            )

        return normalize_zip_path(target)


def load_raw_sheet_cells(
    xlsx_bytes: bytes,
) -> dict[str, dict]:
    """
    Return raw XML information for every physical cell.

    Example:
    {
        "C2": {
            "cell_type": "n",
            "style_id": 3,
            "raw_value": "46205.0"
        }
    }
    """

    worksheet_path = (
        get_first_worksheet_xml_path(
            xlsx_bytes
        )
    )

    cells: dict[str, dict] = {}

    with zipfile.ZipFile(
        io.BytesIO(xlsx_bytes)
    ) as zf:

        xml_bytes = zf.read(
            worksheet_path
        )

        root = ET.fromstring(
            xml_bytes
        )

        for cell in root.findall(
            ".//main:c",
            NS,
        ):

            coordinate = cell.attrib.get("r")

            if not coordinate:
                continue

            cell_type = cell.attrib.get("t")

            style_id_raw = cell.attrib.get("s")

            style_id = None

            if style_id_raw is not None:
                try:
                    style_id = int(style_id_raw)
                except ValueError:
                    style_id = None

            value_node = cell.find(
                "main:v",
                NS,
            )

            raw_value = (
                value_node.text
                if value_node is not None
                else None
            )

            cells[coordinate] = {
                "cell_type": cell_type,
                "style_id": style_id,
                "raw_value": raw_value,
            }

    return cells


# ============================================================
# EXCEL SERIAL ANALYSIS
# ============================================================

def parse_raw_numeric(
    raw_value,
) -> Optional[float]:

    if raw_value is None:
        return None

    try:
        return float(raw_value)
    except (TypeError, ValueError):
        return None


def excel_serial_as_datetime(
    serial_value: Optional[float],
    epoch,
):

    if serial_value is None:
        return None

    try:
        return from_excel(
            serial_value,
            epoch=epoch,
        )
    except Exception:
        return None


def candidate_day_month_recovery(
    interpreted_value,
) -> Optional[float]:
    """
    IMPORTANT:
    This is NOT accepted as the final recovery rule.

    It only computes the hypothesis:

        day + month / 10

    Example:
        datetime(..., month=7, day=2)
        -> 2.7

    The script prints this only for comparison.
    """

    if not isinstance(
        interpreted_value,
        (datetime, date),
    ):
        return None

    return (
        interpreted_value.day
        + interpreted_value.month / 10.0
    )


# ============================================================
# INSPECT SINGLE XLSX
# ============================================================

def inspect_xlsx(
    object_name: str,
):

    report.section(
        f"FILE: {Path(object_name).name}"
    )

    report.write(
        f"Object : {object_name}"
    )

    xlsx_bytes = download_object(
        object_name
    )

    report.write(
        f"Size   : {len(xlsx_bytes)} bytes"
    )

    raw_cells = load_raw_sheet_cells(
        xlsx_bytes
    )

    workbook = load_workbook(
        io.BytesIO(xlsx_bytes),
        data_only=True,
        read_only=True,
    )

    try:

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        report.write(
            f"Sheet  : {worksheet.title}"
        )

        report.write(
            f"Rows   : {worksheet.max_row}"
        )

        report.write(
            f"Cols   : {worksheet.max_column}"
        )

        report.write(
            f"Excel epoch : {workbook.epoch}"
        )

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        headers = [
            worksheet.cell(1, col).value
            for col in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        report.write(
            f"Headers: {headers}"
        )

        # ----------------------------------------------------
        # Statistics
        # ----------------------------------------------------

        height_type_counter = Counter()

        number_format_counter = Counter()

        xml_cell_type_counter = Counter()

        style_counter = Counter()

        datetime_rows = []
        string_numeric_rows = []
        numeric_rows = []

        all_height_rows = []

        # ----------------------------------------------------
        # Read column C = Độ cao (m)
        # ----------------------------------------------------

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            date_cell = worksheet.cell(
                row_number,
                1,
            )

            time_cell = worksheet.cell(
                row_number,
                2,
            )

            height_cell = worksheet.cell(
                row_number,
                3,
            )

            height_value = height_cell.value

            value_type = classify_value(
                height_value
            )

            height_type_counter[
                value_type
            ] += 1

            number_format_counter[
                height_cell.number_format
            ] += 1

            coordinate = (
                height_cell.coordinate
            )

            xml_info = raw_cells.get(
                coordinate,
                {},
            )

            raw_value = xml_info.get(
                "raw_value"
            )

            xml_type = xml_info.get(
                "cell_type"
            )

            style_id = xml_info.get(
                "style_id"
            )

            xml_cell_type_counter[
                str(xml_type)
            ] += 1

            style_counter[
                str(style_id)
            ] += 1

            raw_numeric = parse_raw_numeric(
                raw_value
            )

            serial_datetime = (
                excel_serial_as_datetime(
                    raw_numeric,
                    workbook.epoch,
                )
            )

            recovery_candidate = (
                candidate_day_month_recovery(
                    height_value
                )
            )

            record = {
                "row": row_number,
                "coordinate": coordinate,
                "date": date_cell.value,
                "time": time_cell.value,
                "openpyxl_value": height_value,
                "openpyxl_type": value_type,
                "number_format":
                    height_cell.number_format,
                "xml_type": xml_type,
                "style_id": style_id,
                "raw_xml_value": raw_value,
                "raw_numeric": raw_numeric,
                "serial_datetime":
                    serial_datetime,
                "candidate_recovery":
                    recovery_candidate,
            }

            all_height_rows.append(
                record
            )

            if value_type == "DATETIME":
                datetime_rows.append(
                    record
                )

            elif value_type == "STRING_NUMERIC":
                string_numeric_rows.append(
                    record
                )

            elif value_type == "NUMERIC":
                numeric_rows.append(
                    record
                )

        # ----------------------------------------------------
        # Summary
        # ----------------------------------------------------

        report.subsection(
            "HEIGHT TYPE SUMMARY"
        )

        for key, count in (
            height_type_counter
            .most_common()
        ):
            report.write(
                f"{key:<20}: {count}"
            )

        report.subsection(
            "NUMBER FORMAT SUMMARY"
        )

        for key, count in (
            number_format_counter
            .most_common()
        ):
            report.write(
                f"{key!r:<30}: {count}"
            )

        report.subsection(
            "RAW XML CELL TYPE SUMMARY"
        )

        for key, count in (
            xml_cell_type_counter
            .most_common()
        ):
            report.write(
                f"{key:<20}: {count}"
            )

        report.subsection(
            "STYLE ID SUMMARY"
        )

        for key, count in (
            style_counter
            .most_common()
        ):
            report.write(
                f"{key:<20}: {count}"
            )

        # ----------------------------------------------------
        # DATETIME examples
        # ----------------------------------------------------

        report.subsection(
            "DATETIME HEIGHT CELLS - FIRST 30"
        )

        if not datetime_rows:

            report.write(
                "No DATETIME height cells."
            )

        else:

            for record in datetime_rows[:30]:

                report.write(
                    f"Row {record['row']}"
                )

                report.write(
                    f"  Cell             : "
                    f"{record['coordinate']}"
                )

                report.write(
                    f"  Date             : "
                    f"{record['date']!r}"
                )

                report.write(
                    f"  Time             : "
                    f"{record['time']!r}"
                )

                report.write(
                    f"  OpenPyXL value   : "
                    f"{record['openpyxl_value']!r}"
                )

                report.write(
                    f"  OpenPyXL type    : "
                    f"{record['openpyxl_type']}"
                )

                report.write(
                    f"  Number format    : "
                    f"{record['number_format']!r}"
                )

                report.write(
                    f"  XML cell type    : "
                    f"{record['xml_type']!r}"
                )

                report.write(
                    f"  Style ID         : "
                    f"{record['style_id']!r}"
                )

                report.write(
                    f"  RAW XML <v>      : "
                    f"{record['raw_xml_value']!r}"
                )

                report.write(
                    f"  RAW numeric      : "
                    f"{record['raw_numeric']!r}"
                )

                report.write(
                    f"  Serial -> date   : "
                    f"{record['serial_datetime']!r}"
                )

                report.write(
                    f"  Candidate d+m/10 : "
                    f"{record['candidate_recovery']!r}"
                )

                report.write()

        # ----------------------------------------------------
        # Boundary context
        # ----------------------------------------------------

        report.subsection(
            "TRANSITION / BOUNDARY EXAMPLES"
        )

        transitions = []

        previous_type = None

        for index, record in enumerate(
            all_height_rows
        ):

            current_type = (
                record["openpyxl_type"]
            )

            if (
                previous_type is not None
                and current_type != previous_type
            ):
                transitions.append(index)

            previous_type = current_type

        if not transitions:

            report.write(
                "No type transitions detected."
            )

        else:

            for transition_number, index in enumerate(
                transitions[:10],
                start=1,
            ):

                report.write(
                    f"Transition "
                    f"{transition_number}:"
                )

                start = max(
                    0,
                    index - 2,
                )

                end = min(
                    len(all_height_rows),
                    index + 3,
                )

                for record in (
                    all_height_rows[start:end]
                ):

                    report.write(
                        "  "
                        f"row={record['row']:<5} "
                        f"cell={record['coordinate']:<6} "
                        f"type={record['openpyxl_type']:<15} "
                        f"value="
                        f"{record['openpyxl_value']!r:<30} "
                        f"format="
                        f"{record['number_format']!r:<10} "
                        f"raw_v="
                        f"{record['raw_xml_value']!r}"
                    )

                report.write()

        # ----------------------------------------------------
        # Numeric/string examples for comparison
        # ----------------------------------------------------

        report.subsection(
            "NORMAL NUMERIC / STRING NUMERIC EXAMPLES"
        )

        normal_examples = (
            numeric_rows[:10]
            + string_numeric_rows[:10]
        )

        if not normal_examples:

            report.write(
                "No normal numeric examples."
            )

        else:

            for record in normal_examples:

                report.write(
                    f"row={record['row']:<5} "
                    f"cell={record['coordinate']:<6} "
                    f"type={record['openpyxl_type']:<15} "
                    f"value={record['openpyxl_value']!r:<15} "
                    f"format={record['number_format']!r:<12} "
                    f"xml_type={record['xml_type']!r:<6} "
                    f"style={record['style_id']!r:<5} "
                    f"raw_v={record['raw_xml_value']!r}"
                )

        # ----------------------------------------------------
        # File conclusion
        # ----------------------------------------------------

        report.subsection(
            "FILE SUMMARY"
        )

        report.write(
            f"Total height rows       : "
            f"{len(all_height_rows)}"
        )

        report.write(
            f"DATETIME height rows    : "
            f"{len(datetime_rows)}"
        )

        report.write(
            f"NUMERIC height rows     : "
            f"{len(numeric_rows)}"
        )

        report.write(
            f"STRING_NUMERIC rows     : "
            f"{len(string_numeric_rows)}"
        )

    finally:

        workbook.close()


# ============================================================
# MAIN
# ============================================================

def main():

    report.section(
        "TIDE 2021 RAW EXCEL XML INSPECTION"
    )

    report.write(
        f"MinIO endpoint : {MINIO_ENDPOINT}"
    )

    report.write(
        f"Bucket         : {MINIO_BUCKET}"
    )

    report.write(
        f"Prefix         : {BRONZE_PREFIX}"
    )

    # --------------------------------------------------------
    # Discover 2021 XLSX files
    # --------------------------------------------------------

    objects = list(
        minio_client.list_objects(
            MINIO_BUCKET,
            prefix=BRONZE_PREFIX,
            recursive=True,
        )
    )

    tide_2021_objects = []

    for obj in objects:

        object_name = obj.object_name

        if object_name.endswith("/"):
            continue

        if not object_name.lower().endswith(
            ".xlsx"
        ):
            continue

        year = extract_year_from_filename(
            object_name
        )

        if year != 2021:
            continue

        tide_2021_objects.append(
            object_name
        )

    tide_2021_objects.sort(
        key=lambda name: (
            extract_month_from_filename(name)
            or 99,
            name,
        )
    )

    report.write(
        f"2021 XLSX files found: "
        f"{len(tide_2021_objects)}"
    )

    for object_name in tide_2021_objects:

        report.write(
            f"  - {object_name}"
        )

    if not tide_2021_objects:

        report.write()
        report.write(
            "No Tide 2021 XLSX files found."
        )

        report.save()

        return

    # --------------------------------------------------------
    # Inspect all 2021 XLSX
    # --------------------------------------------------------

    success_count = 0
    failure_count = 0

    for index, object_name in enumerate(
        tide_2021_objects,
        start=1,
    ):

        report.section(
            f"PROCESSING {index}/"
            f"{len(tide_2021_objects)}"
        )

        try:

            inspect_xlsx(
                object_name
            )

            success_count += 1

        except Exception as exc:

            failure_count += 1

            report.write(
                f"[ERROR] {object_name}"
            )

            report.write(
                f"{type(exc).__name__}: "
                f"{exc}"
            )

    # --------------------------------------------------------
    # Final summary
    # --------------------------------------------------------

    report.section(
        "FINAL SUMMARY"
    )

    report.write(
        f"Files discovered : "
        f"{len(tide_2021_objects)}"
    )

    report.write(
        f"Files inspected  : "
        f"{success_count}"
    )

    report.write(
        f"Files failed     : "
        f"{failure_count}"
    )

    report.write()
    report.write(
        "IMPORTANT:"
    )

    report.write(
        "Candidate d+m/10 is printed only "
        "as a hypothesis."
    )

    report.write(
        "Do NOT use it as the Silver recovery "
        "rule until the XML evidence has been "
        "reviewed."
    )

    report.section(
        "REPORT SAVED"
    )

    report.write(
        f"File: {REPORT_PATH}"
    )

    report.save()


if __name__ == "__main__":
    main()