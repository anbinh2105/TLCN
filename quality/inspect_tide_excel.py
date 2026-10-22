from pathlib import Path
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILES = [
    PROJECT_ROOT / "data/raw/tide/thuy_trieu_thang1_2021.xlsx",
    PROJECT_ROOT / "data/raw/tide/thuy_trieu_thang4_2021.xlsx",
]


def inspect_file(file_path: Path):
    print("\n" + "=" * 80)
    print(f"FILE: {file_path.name}")
    print("=" * 80)

    wb = load_workbook(
        file_path,
        data_only=False,
        read_only=False,
    )

    ws = wb.active

    print("\nHeaders:")
    for cell in ws[1]:
        print(
            f"{cell.coordinate}: "
            f"value={repr(cell.value)}, "
            f"type={cell.data_type}, "
            f"format={cell.number_format}"
        )

    print("\nFirst 20 data rows:")

    for row in range(2, min(ws.max_row, 21) + 1):

        print(f"\nROW {row}")

        for col in range(1, 4):

            cell = ws.cell(
                row=row,
                column=col,
            )

            print(
                f"  {cell.coordinate}: "
                f"value={repr(cell.value)}, "
                f"python_type={type(cell.value).__name__}, "
                f"data_type={cell.data_type}, "
                f"is_date={cell.is_date}, "
                f"number_format={repr(cell.number_format)}"
            )

    wb.close()


def main():

    for file_path in FILES:

        if not file_path.exists():
            print(f"NOT FOUND: {file_path}")
            continue

        inspect_file(file_path)


if __name__ == "__main__":
    main()