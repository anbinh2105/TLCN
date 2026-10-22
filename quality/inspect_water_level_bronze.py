import io
import os

from minio import Minio
from openpyxl import load_workbook
from dotenv import load_dotenv


load_dotenv()


# ============================================================
# MinIO configuration
# ============================================================

MINIO_ENDPOINT = os.getenv(
    "MINIO_ENDPOINT",
    "localhost:9000",
)

MINIO_ACCESS_KEY = os.getenv(
    "MINIO_ACCESS_KEY",
    "admin",
)

MINIO_SECRET_KEY = os.getenv(
    "MINIO_SECRET_KEY",
    "admin12345",
)

MINIO_BUCKET = os.getenv(
    "MINIO_BUCKET",
    "flood-lakehouse",
)

MINIO_SECURE = (
    os.getenv("MINIO_SECURE", "false")
    .strip()
    .lower()
    == "true"
)

PREFIX = "bronze/water_level/"


# ============================================================
# Connect MinIO
# ============================================================

client = Minio(
    MINIO_ENDPOINT,
    access_key=MINIO_ACCESS_KEY,
    secret_key=MINIO_SECRET_KEY,
    secure=MINIO_SECURE,
)


# ============================================================
# Find XLSX objects
# ============================================================

objects = list(
    client.list_objects(
        MINIO_BUCKET,
        prefix=PREFIX,
        recursive=True,
    )
)

xlsx_objects = [
    obj
    for obj in objects
    if obj.object_name.lower().endswith(".xlsx")
]


print("=" * 80)
print("WATER LEVEL BRONZE INSPECTION")
print("=" * 80)

print(f"Bucket       : {MINIO_BUCKET}")
print(f"Prefix       : {PREFIX}")
print(f"XLSX objects : {len(xlsx_objects)}")
print()


if not xlsx_objects:
    raise RuntimeError(
        "No XLSX Water Level objects found in Bronze."
    )


# ============================================================
# Print object list
# ============================================================

print("BRONZE XLSX OBJECTS")
print("-" * 80)

for i, obj in enumerate(xlsx_objects, start=1):
    print(
        f"{i:02d}. "
        f"{obj.object_name} "
        f"({obj.size} bytes)"
    )


# ============================================================
# Inspect first XLSX file
# ============================================================

target = xlsx_objects[0]

print()
print("=" * 80)
print("INSPECTING FIRST XLSX")
print("=" * 80)

print(f"Object: {target.object_name}")


response = client.get_object(
    MINIO_BUCKET,
    target.object_name,
)

try:
    file_bytes = response.read()
finally:
    response.close()
    response.release_conn()


# ============================================================
# Load workbook
# ============================================================

workbook = load_workbook(
    io.BytesIO(file_bytes),
    data_only=True,
    read_only=False,
)


print()
print("Sheet names:")
for sheet_name in workbook.sheetnames:
    print(f"  - {sheet_name}")


# ============================================================
# Inspect each sheet
# ============================================================

for sheet_name in workbook.sheetnames:

    ws = workbook[sheet_name]

    print()
    print("=" * 80)
    print(f"SHEET: {sheet_name}")
    print("=" * 80)

    print(
        f"Rows    : {ws.max_row}"
    )

    print(
        f"Columns : {ws.max_column}"
    )

    print()

    max_rows = min(
        ws.max_row,
        20,
    )

    max_cols = min(
        ws.max_column,
        15,
    )

    for row_idx in range(
        1,
        max_rows + 1,
    ):

        print(
            f"ROW {row_idx}"
        )

        for col_idx in range(
            1,
            max_cols + 1,
        ):

            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            print(
                f"  {cell.coordinate:<5} "
                f"value={repr(cell.value):<30} "
                f"type={cell.data_type:<3} "
                f"is_date={cell.is_date!s:<5} "
                f"format={repr(cell.number_format)}"
            )

        print("-" * 80)


workbook.close()


print()
print("=" * 80)
print("WATER LEVEL BRONZE INSPECTION COMPLETE")
print("=" * 80)