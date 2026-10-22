import csv
import io
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
import sys
from pathlib import Path

from dotenv import load_dotenv
from minio import Minio
from openpyxl import load_workbook

import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


# ============================================================
# CONFIG
# ============================================================

load_dotenv()

MINIO_ENDPOINT = os.getenv("MINIO_ENDPOINT", "localhost:9000")
MINIO_ACCESS_KEY = os.getenv("MINIO_ACCESS_KEY", "admin")
MINIO_SECRET_KEY = os.getenv("MINIO_SECRET_KEY", "admin12345")
MINIO_BUCKET = os.getenv("MINIO_BUCKET", "flood-lakehouse")
MINIO_SECURE = os.getenv("MINIO_SECURE", "false").lower() == "true"

BRONZE_PREFIX = "bronze/tide/"


# ============================================================
# MINIO
# ============================================================

client = Minio(
    MINIO_ENDPOINT,
    access_key=MINIO_ACCESS_KEY,
    secret_key=MINIO_SECRET_KEY,
    secure=MINIO_SECURE,
)


def get_object_bytes(object_name):
    response = client.get_object(
        MINIO_BUCKET,
        object_name,
    )

    try:
        return response.read()
    finally:
        response.close()
        response.release_conn()


# ============================================================
# HELPERS
# ============================================================

def separator():
    print("=" * 100)


def extract_year_from_name(name):
    """
    Try to find 2021, 2022, ..., 2025 from filename/path.
    """

    match = re.search(
        r"(2021|2022|2023|2024|2025)",
        name,
    )

    if match:
        return int(match.group(1))

    return None


def normalize_header(value):
    if value is None:
        return None

    return str(value).strip()


def classify_excel_height(cell):
    """
    IMPORTANT:
    This function DOES NOT recover/fix the value.

    It only classifies how OpenPyXL sees the cell.
    """

    value = cell.value

    if value is None:
        return "NULL"

    if isinstance(value, datetime):
        return "DATETIME"

    if isinstance(value, date):
        return "DATE"

    if isinstance(value, time):
        return "TIME"

    if isinstance(value, bool):
        return "BOOLEAN"

    if isinstance(value, (int, float)):
        return "NUMERIC"

    if isinstance(value, str):
        text = value.strip()

        if text == "":
            return "EMPTY_STRING"

        try:
            float(text)
            return "STRING_NUMERIC"
        except ValueError:
            return "STRING_OTHER"

    return type(value).__name__.upper()


def classify_csv_height(value):
    if value is None:
        return "NULL"

    text = str(value).strip()

    if text == "":
        return "EMPTY_STRING"

    try:
        float(text)
        return "NUMERIC_STRING"
    except ValueError:
        return "STRING_OTHER"


# ============================================================
# XLSX INSPECTION
# ============================================================

def inspect_xlsx(object_name, file_bytes):
    print()
    separator()
    print("XLSX FILE")
    separator()

    print("Object :", object_name)
    print("Size   :", len(file_bytes), "bytes")

    wb = load_workbook(
        io.BytesIO(file_bytes),
        data_only=True,
        read_only=False,
    )

    print("Sheets :", wb.sheetnames)

    # Tide files are expected to have data in first sheet.
    ws = wb[wb.sheetnames[0]]

    print(
        "Active data sheet:",
        ws.title,
    )

    print(
        "Dimensions:",
        ws.max_row,
        "rows x",
        ws.max_column,
        "columns",
    )

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        normalize_header(
            ws.cell(
                row=1,
                column=column,
            ).value
        )
        for column in range(
            1,
            ws.max_column + 1,
        )
    ]

    print()
    print("Headers:")

    for index, header in enumerate(
        headers,
        start=1,
    ):
        print(
            f"  Column {index}: "
            f"{repr(header)}"
        )

    # --------------------------------------------------------
    # First 10 records
    # --------------------------------------------------------

    print()
    print("First 10 data rows:")

    for row_number in range(
        2,
        min(ws.max_row + 1, 12),
    ):
        values = [
            ws.cell(
                row=row_number,
                column=column,
            ).value
            for column in range(
                1,
                ws.max_column + 1,
            )
        ]

        print(
            f"  Row {row_number}: {values}"
        )

    # --------------------------------------------------------
    # Assume first 3 columns:
    #
    # A = Ngày
    # B = Giờ
    # C = Độ cao (m)
    #
    # We are inspecting only.
    # --------------------------------------------------------

    if ws.max_column < 3:
        print(
            "[WARNING] Less than 3 columns. "
            "Cannot inspect Tide structure."
        )

        wb.close()

        return {
            "type": "xlsx",
            "rows": max(ws.max_row - 1, 0),
            "headers": headers,
        }

    height_types = Counter()
    height_formats = Counter()

    numeric_values = []

    suspicious_examples = defaultdict(list)

    date_type_counter = Counter()
    time_type_counter = Counter()

    total_rows = 0

    # --------------------------------------------------------
    # Scan all rows
    # --------------------------------------------------------

    for row_number in range(
        2,
        ws.max_row + 1,
    ):
        date_cell = ws.cell(
            row=row_number,
            column=1,
        )

        time_cell = ws.cell(
            row=row_number,
            column=2,
        )

        height_cell = ws.cell(
            row=row_number,
            column=3,
        )

        if (
            date_cell.value is None
            and time_cell.value is None
            and height_cell.value is None
        ):
            continue

        total_rows += 1

        # Date / Time representations
        date_type_counter[
            type(date_cell.value).__name__
        ] += 1

        time_type_counter[
            type(time_cell.value).__name__
        ] += 1

        # Height
        category = classify_excel_height(
            height_cell
        )

        height_types[category] += 1

        height_formats[
            str(height_cell.number_format)
        ] += 1

        if category == "NUMERIC":
            numeric_values.append(
                float(height_cell.value)
            )

        if category in {
            "DATETIME",
            "DATE",
            "TIME",
            "STRING_NUMERIC",
            "STRING_OTHER",
        }:
            if len(
                suspicious_examples[category]
            ) < 15:
                suspicious_examples[
                    category
                ].append(
                    {
                        "row": row_number,
                        "date": date_cell.value,
                        "time": time_cell.value,
                        "height_value": (
                            height_cell.value
                        ),
                        "height_data_type": (
                            height_cell.data_type
                        ),
                        "height_is_date": (
                            height_cell.is_date
                        ),
                        "height_number_format": (
                            height_cell.number_format
                        ),
                    }
                )

    print()
    print("Total non-empty data rows:", total_rows)

    print()
    print("Date column Python types:")

    for key, count in sorted(
        date_type_counter.items()
    ):
        print(
            f"  {key:<20}: {count}"
        )

    print()
    print("Time column Python types:")

    for key, count in sorted(
        time_type_counter.items()
    ):
        print(
            f"  {key:<20}: {count}"
        )

    print()
    print("HEIGHT TYPE DISTRIBUTION:")

    for key, count in sorted(
        height_types.items()
    ):
        print(
            f"  {key:<20}: {count}"
        )

    print()
    print("HEIGHT NUMBER FORMATS:")

    for key, count in height_formats.most_common():
        print(
            f"  {repr(key):<30}: {count}"
        )

    if numeric_values:
        print()
        print("Numeric height statistics:")

        print(
            "  min:",
            min(numeric_values),
        )

        print(
            "  max:",
            max(numeric_values),
        )

        print(
            "  count:",
            len(numeric_values),
        )

    # --------------------------------------------------------
    # Suspicious examples
    # --------------------------------------------------------

    print()
    print("SUSPICIOUS / SPECIAL HEIGHT EXAMPLES")

    for category, examples in (
        suspicious_examples.items()
    ):
        print()
        print(
            f"[{category}] "
            f"showing {len(examples)} examples"
        )

        for example in examples:
            print(
                "  "
                f"row={example['row']} | "
                f"date={repr(example['date'])} | "
                f"time={repr(example['time'])} | "
                f"height={repr(example['height_value'])} | "
                f"data_type={example['height_data_type']} | "
                f"is_date={example['height_is_date']} | "
                f"format={repr(example['height_number_format'])}"
            )

    wb.close()

    return {
        "type": "xlsx",
        "rows": total_rows,
        "headers": headers,
        "height_types": dict(height_types),
        "height_formats": dict(height_formats),
    }


# ============================================================
# CSV INSPECTION
# ============================================================

def detect_encoding(file_bytes):
    """
    Try common encodings used by Vietnamese CSV/Excel exports.
    """

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1258",
        "latin1",
    ]

    for encoding in encodings:
        try:
            text = file_bytes.decode(
                encoding
            )

            return encoding, text

        except UnicodeDecodeError:
            continue

    raise UnicodeDecodeError(
        "unknown",
        b"",
        0,
        1,
        "Could not decode CSV",
    )


def detect_delimiter(text):
    sample = text[:10000]

    try:
        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=[
                ",",
                ";",
                "\t",
                "|",
            ],
        )

        return dialect.delimiter

    except csv.Error:
        return ","


def inspect_csv(object_name, file_bytes):
    print()
    separator()
    print("CSV FILE")
    separator()

    print("Object :", object_name)
    print("Size   :", len(file_bytes), "bytes")

    encoding, text = detect_encoding(
        file_bytes
    )

    delimiter = detect_delimiter(
        text
    )

    print("Encoding :", encoding)
    print(
        "Delimiter:",
        repr(delimiter),
    )

    stream = io.StringIO(text)

    reader = csv.reader(
        stream,
        delimiter=delimiter,
    )

    rows = list(reader)

    if not rows:
        print("[WARNING] Empty CSV.")

        return {
            "type": "csv",
            "rows": 0,
            "headers": [],
        }

    headers = [
        normalize_header(value)
        for value in rows[0]
    ]

    print()
    print("Headers:")

    for index, header in enumerate(
        headers,
        start=1,
    ):
        print(
            f"  Column {index}: "
            f"{repr(header)}"
        )

    print()
    print("First 10 data rows:")

    for row_number, row in enumerate(
        rows[1:11],
        start=2,
    ):
        print(
            f"  Row {row_number}: {row}"
        )

    height_types = Counter()

    total_rows = 0

    for row in rows[1:]:
        if not any(
            str(value).strip()
            for value in row
        ):
            continue

        total_rows += 1

        if len(row) >= 3:
            category = classify_csv_height(
                row[2]
            )

            height_types[
                category
            ] += 1

    print()
    print(
        "Total non-empty data rows:",
        total_rows,
    )

    print()
    print("HEIGHT TYPE DISTRIBUTION:")

    for key, count in sorted(
        height_types.items()
    ):
        print(
            f"  {key:<20}: {count}"
        )

    return {
        "type": "csv",
        "rows": total_rows,
        "headers": headers,
        "encoding": encoding,
        "delimiter": delimiter,
        "height_types": dict(height_types),
    }

# ============================================================
# UTF-8 OUTPUT / LOG FILE
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

REPORT_DIR = PROJECT_ROOT / "reports" / "quality"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

REPORT_FILE = REPORT_DIR / "tide_bronze_inspection.txt"


class Tee:
    """
    Ghi output đồng thời ra terminal và file UTF-8.
    """

    def __init__(self, *streams):
        self.streams = streams

    def write(self, data):
        for stream in self.streams:
            stream.write(data)
            stream.flush()

    def flush(self):
        for stream in self.streams:
            stream.flush()

# ============================================================
# MAIN
# ============================================================

def main():
    separator()
    print("CAN THO FLOOD LAKEHOUSE")
    print("BRONZE TIDE FULL INSPECTION")
    separator()

    print()
    print("Bucket :", MINIO_BUCKET)
    print("Prefix :", BRONZE_PREFIX)

    objects = [
        obj
        for obj in client.list_objects(
            MINIO_BUCKET,
            prefix=BRONZE_PREFIX,
            recursive=True,
        )
        if obj.object_name.lower().endswith(
            (
                ".xlsx",
                ".xlsm",
                ".csv",
            )
        )
    ]

    objects.sort(
        key=lambda obj: obj.object_name
    )

    print()
    print(
        "Total Tide objects:",
        len(objects),
    )

    if not objects:
        print(
            "[ERROR] No Tide files found."
        )

        return

    # --------------------------------------------------------
    # Inventory
    # --------------------------------------------------------

    print()
    separator()
    print("1. TIDE BRONZE INVENTORY")
    separator()

    year_inventory = Counter()

    for index, obj in enumerate(
        objects,
        start=1,
    ):
        year = extract_year_from_name(
            obj.object_name
        )

        extension = os.path.splitext(
            obj.object_name
        )[1].lower()

        year_inventory[
            (year, extension)
        ] += 1

        print(
            f"{index:02d}. "
            f"{obj.object_name}"
        )

        print(
            f"    size={obj.size} bytes | "
            f"year={year} | "
            f"type={extension}"
        )

    # --------------------------------------------------------
    # Inspect each object
    # --------------------------------------------------------

    results = []

    print()
    separator()
    print("2. PER-FILE INSPECTION")
    separator()

    for index, obj in enumerate(
        objects,
        start=1,
    ):
        print()
        print("#" * 100)
        print(
            f"FILE {index}/{len(objects)}"
        )
        print("#" * 100)

        object_name = obj.object_name

        extension = os.path.splitext(
            object_name
        )[1].lower()

        year = extract_year_from_name(
            object_name
        )

        file_bytes = get_object_bytes(
            object_name
        )

        try:
            if extension in {
                ".xlsx",
                ".xlsm",
            }:
                result = inspect_xlsx(
                    object_name,
                    file_bytes,
                )

            elif extension == ".csv":
                result = inspect_csv(
                    object_name,
                    file_bytes,
                )

            else:
                continue

            result["object_name"] = (
                object_name
            )

            result["year"] = year

            results.append(
                result
            )

        except Exception as exc:
            print()
            print(
                "[ERROR] Failed to inspect:",
                object_name,
            )

            print(
                type(exc).__name__,
                ":",
                exc,
            )

            results.append(
                {
                    "object_name": object_name,
                    "year": year,
                    "type": extension,
                    "error": str(exc),
                }
            )

    # --------------------------------------------------------
    # Final summary
    # --------------------------------------------------------

    print()
    separator()
    print("3. FINAL TIDE SUMMARY")
    separator()

    print()
    print("Inventory by year/type:")

    for (
        year,
        extension,
    ), count in sorted(
        year_inventory.items(),
        key=lambda x: (
            x[0][0] or 0,
            x[0][1],
        ),
    ):
        print(
            f"  year={year} | "
            f"type={extension} | "
            f"files={count}"
        )

    print()
    print("Schema summary:")

    for result in results:
        print()
        print(
            "Object:",
            result.get(
                "object_name"
            ),
        )

        print(
            "  year    :",
            result.get("year"),
        )

        print(
            "  type    :",
            result.get("type"),
        )

        print(
            "  rows    :",
            result.get("rows"),
        )

        print(
            "  headers :",
            result.get("headers"),
        )

        if result.get(
            "height_types"
        ):
            print(
                "  height types:",
                result[
                    "height_types"
                ],
            )

        if result.get(
            "height_formats"
        ):
            print(
                "  height formats:",
                result[
                    "height_formats"
                ],
            )

        if result.get("error"):
            print(
                "  ERROR:",
                result["error"],
            )

    print()
    separator()
    print("TIDE BRONZE INSPECTION COMPLETE")
    separator()


if __name__ == "__main__":

    # Windows terminal -> UTF-8
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(
                encoding="utf-8",
                errors="replace",
            )
            sys.stderr.reconfigure(
                encoding="utf-8",
                errors="replace",
            )
        except Exception:
            pass

    # Python tự tạo report UTF-8
    with open(
        REPORT_FILE,
        "w",
        encoding="utf-8",
        newline="",
    ) as report_file:

        original_stdout = sys.stdout
        original_stderr = sys.stderr

        sys.stdout = Tee(
            original_stdout,
            report_file,
        )

        sys.stderr = Tee(
            original_stderr,
            report_file,
        )

        try:
            main()

            print()
            print("=" * 100)
            print("REPORT SAVED")
            print("=" * 100)
            print(f"File: {REPORT_FILE}")

        finally:
            sys.stdout = original_stdout
            sys.stderr = original_stderr