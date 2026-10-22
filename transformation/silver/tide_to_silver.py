from __future__ import annotations

import csv
import io
import math
import os
import re
from datetime import date, datetime, time
from typing import Any

from minio import Minio
from openpyxl import load_workbook
from pyspark.sql import DataFrame
from pyspark.sql import functions as F
from pyspark.sql import types as T

from transformation.silver.common import (
    MINIO_ACCESS_KEY,
    MINIO_BUCKET,
    MINIO_ENDPOINT,
    MINIO_SECRET_KEY,
    MINIO_SECURE,
    QUALITY_INVALID,
    QUALITY_MISSING,
    QUALITY_OBSERVED,
    QUALITY_RECOVERED,
    add_silver_metadata,
    create_spark_session,
    ensure_silver_namespace,
    get_silver_table,
)


SOURCE = "tide"
BRONZE_PREFIX = f"bronze/{SOURCE}/"

# ---------------------------------------------------------------------
# Silver schema
# ---------------------------------------------------------------------

TIDE_SCHEMA = T.StructType(
    [
        T.StructField("date_raw", T.StringType(), True),
        T.StructField("time_raw", T.StringType(), True),
        T.StructField("height_raw", T.StringType(), True),
        T.StructField("tide_height_m", T.DoubleType(), True),
        T.StructField("height_parse_status", T.StringType(), False),
        T.StructField("recovery_method", T.StringType(), True),
        T.StructField("_bronze_source_file", T.StringType(), False),
    ]
)


# ---------------------------------------------------------------------
# MinIO
# ---------------------------------------------------------------------

def get_minio_client() -> Minio:
    return Minio(
        MINIO_ENDPOINT,
        access_key=MINIO_ACCESS_KEY,
        secret_key=MINIO_SECRET_KEY,
        secure=MINIO_SECURE,
    )


def list_tide_objects(client: Minio) -> list[str]:
    objects = client.list_objects(
        MINIO_BUCKET,
        prefix=BRONZE_PREFIX,
        recursive=True,
    )

    result = []

    for obj in objects:
        name = obj.object_name

        if name.lower().endswith((".xlsx", ".csv")):
            result.append(name)

    return sorted(result)


def download_object(client: Minio, object_name: str) -> bytes:
    response = client.get_object(MINIO_BUCKET, object_name)

    try:
        return response.read()
    finally:
        response.close()
        response.release_conn()


# ---------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------

def normalize_text(value: Any) -> str | None:
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    return value


def numeric_or_none(value: Any) -> float | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        number = float(value)

        if math.isnan(number):
            return None

        return number

    text = str(value).strip()

    if not text:
        return None

    # decimal comma support if encountered
    text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def stringify_date(value: Any) -> str | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return normalize_text(value)


def stringify_time(value: Any) -> str | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")

    if isinstance(value, time):
        return value.strftime("%H:%M:%S")

    return normalize_text(value)


# ---------------------------------------------------------------------
# Excel height recovery
# ---------------------------------------------------------------------

def recover_excel_date_height(
    value: Any,
    number_format: str | None,
) -> tuple[float | None, str, str | None]:
    """
    Return:
        tide_height_m
        height_parse_status
        recovery_method

    Rules established from Tide 2021 forensic inspection:

    1. Numeric / numeric-string:
       direct value -> OBSERVED

    2. OpenPyXL DATE/DATETIME with Excel format d\\.m:
       intended height = day + month / 10
       -> RECOVERED

    3. Otherwise:
       invalid / missing.
    """

    if value is None:
        return None, QUALITY_MISSING, None

    # ---------------------------------------------------------------
    # Important:
    # datetime must be checked BEFORE date because datetime is a
    # subclass of date in Python.
    # ---------------------------------------------------------------

    if isinstance(value, datetime):
        fmt = (number_format or "").lower().replace(" ", "")

        if "d" in fmt and "m" in fmt:
            recovered = float(value.day) + float(value.month) / 10.0

            return (
                round(recovered, 1),
                QUALITY_RECOVERED,
                "excel_date_d_m",
            )

        return None, QUALITY_INVALID, None

    if isinstance(value, date):
        fmt = (number_format or "").lower().replace(" ", "")

        if "d" in fmt and "m" in fmt:
            recovered = float(value.day) + float(value.month) / 10.0

            return (
                round(recovered, 1),
                QUALITY_RECOVERED,
                "excel_date_d_m",
            )

        return None, QUALITY_INVALID, None

    # Normal numeric / numeric string
    number = numeric_or_none(value)

    if number is not None:
        return number, QUALITY_OBSERVED, None

    return None, QUALITY_INVALID, None


# ---------------------------------------------------------------------
# XLSX reader
# ---------------------------------------------------------------------

def read_tide_xlsx(
    object_name: str,
    file_bytes: bytes,
) -> list[tuple]:
    workbook = load_workbook(
        io.BytesIO(file_bytes),
        data_only=True,
        read_only=False,
    )

    worksheet = workbook[workbook.sheetnames[0]]

    rows = []

    for row_number in range(2, worksheet.max_row + 1):
        date_cell = worksheet.cell(row=row_number, column=1)
        time_cell = worksheet.cell(row=row_number, column=2)
        height_cell = worksheet.cell(row=row_number, column=3)

        raw_date = date_cell.value
        raw_time = time_cell.value
        raw_height = height_cell.value

        # Completely empty row
        if (
            raw_date is None
            and raw_time is None
            and raw_height is None
        ):
            continue

        tide_height_m, parse_status, recovery_method = (
            recover_excel_date_height(
                raw_height,
                height_cell.number_format,
            )
        )

        rows.append(
            (
                stringify_date(raw_date),
                stringify_time(raw_time),
                normalize_text(raw_height),
                tide_height_m,
                parse_status,
                recovery_method,
                object_name,
            )
        )

    workbook.close()

    return rows


# ---------------------------------------------------------------------
# CSV reader
# ---------------------------------------------------------------------

def detect_encoding(file_bytes: bytes) -> str:
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1258",
        "latin1",
    ]

    for encoding in encodings:
        try:
            file_bytes.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue

    return "latin1"


def read_tide_csv(
    object_name: str,
    file_bytes: bytes,
) -> list[tuple]:
    encoding = detect_encoding(file_bytes)
    text = file_bytes.decode(encoding)

    sample = text[:4096]

    try:
        dialect = csv.Sniffer().sniff(sample)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    reader = csv.reader(
        io.StringIO(text),
        delimiter=delimiter,
    )

    all_rows = list(reader)

    if not all_rows:
        return []

    rows = []

    # First row = header
    for raw_row in all_rows[1:]:
        if not raw_row:
            continue

        if len(raw_row) < 3:
            continue

        raw_date = normalize_text(raw_row[0])
        raw_time = normalize_text(raw_row[1])
        raw_height = normalize_text(raw_row[2])

        if (
            raw_date is None
            and raw_time is None
            and raw_height is None
        ):
            continue

        tide_height_m = numeric_or_none(raw_height)

        if raw_height is None:
            parse_status = QUALITY_MISSING
        elif tide_height_m is None:
            parse_status = QUALITY_INVALID
        else:
            parse_status = QUALITY_OBSERVED

        rows.append(
            (
                raw_date,
                raw_time,
                raw_height,
                tide_height_m,
                parse_status,
                None,
                object_name,
            )
        )

    return rows


# ---------------------------------------------------------------------
# Read all Bronze Tide objects
# ---------------------------------------------------------------------

def read_bronze_tide(spark) -> DataFrame:
    client = get_minio_client()

    object_names = list_tide_objects(client)

    if not object_names:
        raise RuntimeError(
            f"No Tide files found under "
            f"s3a://{MINIO_BUCKET}/{BRONZE_PREFIX}"
        )

    print("=" * 100)
    print("TIDE BRONZE FILES")
    print("=" * 100)

    all_rows = []

    for object_name in object_names:
        print(f"Reading: {object_name}")

        file_bytes = download_object(
            client,
            object_name,
        )

        extension = os.path.splitext(object_name)[1].lower()

        if extension == ".xlsx":
            rows = read_tide_xlsx(
                object_name,
                file_bytes,
            )

        elif extension == ".csv":
            rows = read_tide_csv(
                object_name,
                file_bytes,
            )

        else:
            continue

        print(f"  Records: {len(rows):,}")

        all_rows.extend(rows)

    if not all_rows:
        raise RuntimeError("No Tide records could be read.")

    print("-" * 100)
    print(f"Total Bronze Tide records: {len(all_rows):,}")
    print("=" * 100)

    return spark.createDataFrame(
        all_rows,
        schema=TIDE_SCHEMA,
    )


# ---------------------------------------------------------------------
# Timestamp normalization
# ---------------------------------------------------------------------

def build_timestamp(df: DataFrame) -> DataFrame:
    """
    Build canonical hourly Tide timestamp from date_raw + time_raw.

    Handles:
      Date:
        2021-01-01
        2021/01/01
        01/01/2021
        01-01-2021

      Time:
        00:00:00
        00:00
        0:00
        0
        1
        ...
        23

    Bare integer hours are a representation found in the original
    XLSX source and are standardized to HH:00:00.
    """

    date_str = F.trim(F.col("date_raw").cast("string"))
    time_str = F.trim(F.col("time_raw").cast("string"))

    # ---------------------------------------------------------
    # 1. Normalize DATE to yyyy-MM-dd
    # ---------------------------------------------------------
    normalized_date = (
        F.when(
            date_str.rlike(r"^\d{4}-\d{1,2}-\d{1,2}$"),
            F.regexp_replace(date_str, "/", "-"),
        )
        .when(
            date_str.rlike(r"^\d{4}/\d{1,2}/\d{1,2}$"),
            F.regexp_replace(date_str, "/", "-"),
        )
        .when(
            date_str.rlike(r"^\d{1,2}/\d{1,2}/\d{4}$"),
            F.date_format(
                F.try_to_timestamp(
                    date_str,
                    F.lit("d/M/yyyy"),
                ),
                "yyyy-MM-dd",
            ),
        )
        .when(
            date_str.rlike(r"^\d{1,2}-\d{1,2}-\d{4}$"),
            F.date_format(
                F.try_to_timestamp(
                    date_str,
                    F.lit("d-M-yyyy"),
                ),
                "yyyy-MM-dd",
            ),
        )
    )

    # ---------------------------------------------------------
    # 2. Normalize TIME
    # ---------------------------------------------------------

    # 0 ... 23
    integer_hour = time_str.rlike(
        r"^(?:[0-9]|1[0-9]|2[0-3])$"
    )

    normalized_time = (
        F.when(
            integer_hour,
            F.concat(
                F.lpad(time_str, 2, "0"),
                F.lit(":00:00"),
            ),
        )
        .when(
            time_str.rlike(r"^\d{1,2}:\d{2}$"),
            F.concat(
                time_str,
                F.lit(":00"),
            ),
        )
        .otherwise(time_str)
    )

    # ---------------------------------------------------------
    # 3. Combine DATE + TIME
    # ---------------------------------------------------------
    timestamp_string = F.concat_ws(
        " ",
        normalized_date,
        normalized_time,
    )

    # ---------------------------------------------------------
    # 4. Parse safely
    # Spark 4 ANSI-safe: use try_to_timestamp
    # ---------------------------------------------------------
    timestamp = F.coalesce(
        F.try_to_timestamp(
            timestamp_string,
            F.lit("yyyy-MM-dd HH:mm:ss"),
        ),
        F.try_to_timestamp(
            timestamp_string,
            F.lit("yyyy-MM-dd H:mm:ss"),
        ),
    )

    return df.withColumn("timestamp", timestamp)

# ---------------------------------------------------------------------
# Quality classification
# ---------------------------------------------------------------------

def assign_quality_flag(df: DataFrame) -> DataFrame:
    return df.withColumn(
        "quality_flag",
        F.when(
            F.col("timestamp").isNull(),
            F.lit(QUALITY_INVALID),
        )
        .when(
            F.col("height_parse_status") == QUALITY_INVALID,
            F.lit(QUALITY_INVALID),
        )
        .when(
            F.col("tide_height_m").isNull(),
            F.lit(QUALITY_MISSING),
        )
        .when(
            F.col("height_parse_status") == QUALITY_RECOVERED,
            F.lit(QUALITY_RECOVERED),
        )
        .otherwise(
            F.lit(QUALITY_OBSERVED)
        ),
    )


# ---------------------------------------------------------------------
# Bronze -> Silver
# ---------------------------------------------------------------------

def transform_to_silver(df: DataFrame) -> DataFrame:
    df = build_timestamp(df)

    # -------------------------------------------------------------
    # DEBUG / DATA PROFILING:
    # Show the original Date + Time representation that Spark
    # could not parse.
    # -------------------------------------------------------------
    print()
    print("=" * 100)
    print("UNPARSED TIDE TIMESTAMP PROFILING")
    print("=" * 100)

    unparsed = df.filter(
        F.col("timestamp").isNull()
    )

    print(
        f"Unparsed timestamps: "
        f"{unparsed.count():,}"
    )

    print()
    print("Unparsed by source file:")

    (
        unparsed
        .groupBy("_bronze_source_file")
        .count()
        .orderBy("_bronze_source_file")
        .show(100, truncate=False)
    )

    print()
    print("Sample raw Date / Time values:")

    (
        unparsed
        .select(
            "date_raw",
            "time_raw",
            "_bronze_source_file",
        )
        .show(
            100,
            truncate=False,
        )
    )

    df = assign_quality_flag(df)

    df = add_silver_metadata(
        df,
        source=SOURCE,
        source_file_column="_bronze_source_file",
    )

    return df.select(
        "timestamp",
        "tide_height_m",
        "quality_flag",
        "recovery_method",
        "source",
        "source_file",
        "silver_processed_at",
    )


# ---------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------

def validate_silver(df: DataFrame) -> None:
    print()
    print("=" * 100)
    print("TIDE SILVER VALIDATION")
    print("=" * 100)

    total = df.count()

    print(f"Silver records : {total:,}")

    print()
    print("Quality distribution:")
    df.groupBy("quality_flag").count().orderBy(
        F.desc("count")
    ).show(truncate=False)

    print("Recovery methods:")
    df.groupBy("recovery_method").count().orderBy(
        F.desc("count")
    ).show(truncate=False)

    null_timestamp = df.filter(
        F.col("timestamp").isNull()
    ).count()

    null_height = df.filter(
        F.col("tide_height_m").isNull()
    ).count()

    print(f"Null timestamp : {null_timestamp:,}")
    print(f"Null tide height: {null_height:,}")

    print()
    print("Tide height range:")

    df.select(
        F.min("tide_height_m").alias("min_tide_height_m"),
        F.max("tide_height_m").alias("max_tide_height_m"),
    ).show()

    print()
    print("Duplicate timestamps:")

    duplicate_count = (
        df.groupBy("timestamp")
        .count()
        .filter(
            F.col("timestamp").isNotNull()
            & (F.col("count") > 1)
        )
        .count()
    )

    print(f"Duplicate timestamp groups: {duplicate_count:,}")

    print()
    print("Sample RECOVERED records:")

    df.filter(
        F.col("quality_flag") == QUALITY_RECOVERED
    ).orderBy("timestamp").show(
        30,
        truncate=False,
    )


# ---------------------------------------------------------------------
# Write Iceberg
# ---------------------------------------------------------------------

def write_silver(df: DataFrame) -> None:
    table_name = get_silver_table(SOURCE)

    # ---------------------------------------------------------
    # SILVER QUALITY GATE
    # ---------------------------------------------------------

    print()
    print("=" * 100)
    print("TIDE SILVER QUALITY GATE")
    print("=" * 100)

    critical_null_timestamp = df.filter(
        F.col("timestamp").isNull()
    ).count()

    critical_null_height = df.filter(
        F.col("tide_height_m").isNull()
    ).count()

    critical_invalid = df.filter(
        F.col("quality_flag") == QUALITY_INVALID
    ).count()

    print(f"Null timestamp : {critical_null_timestamp:,}")
    print(f"Null height    : {critical_null_height:,}")
    print(f"INVALID records: {critical_invalid:,}")

    # Critical failures:
    # - timestamp is required for time-series data
    # - INVALID records must not be published to Silver
    #
    # Missing tide height is reported but is not automatically
    # treated as a critical failure because Silver may preserve
    # legitimate MISSING observations.
    if critical_null_timestamp > 0 or critical_invalid > 0:
        raise RuntimeError(
            "\nTIDE SILVER QUALITY GATE FAILED\n"
            f"  Null timestamp : {critical_null_timestamp:,}\n"
            f"  Null height    : {critical_null_height:,}\n"
            f"  INVALID records: {critical_invalid:,}\n"
            "Iceberg table was NOT written."
        )

    print("TIDE SILVER QUALITY GATE: PASS")

    # ---------------------------------------------------------
    # WRITE TO ICEBERG
    # ---------------------------------------------------------

    print()
    print("=" * 100)
    print("WRITE TIDE SILVER")
    print("=" * 100)
    print(f"Target table: {table_name}")

    (
        df.writeTo(table_name)
        .using("iceberg")
        .tableProperty(
            "format-version",
            "2",
        )
        .createOrReplace()
    )

    print("Iceberg write completed.")
# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def main() -> None:
    spark = create_spark_session(
        app_name="Tide-Bronze-to-Silver"
    )

    try:
        ensure_silver_namespace(spark)

        bronze_df = read_bronze_tide(spark)

        silver_df = transform_to_silver(
            bronze_df
        ).cache()

        validate_silver(silver_df)

        write_silver(silver_df)

        print()
        print("=" * 100)
        print("VERIFY ICEBERG TABLE")
        print("=" * 100)

        iceberg_df = spark.table(
            get_silver_table(SOURCE)
        )

        print(
            f"Iceberg records: "
            f"{iceberg_df.count():,}"
        )

        iceberg_df.orderBy(
            "timestamp"
        ).show(
            30,
            truncate=False,
        )

        print()
        print("=" * 100)
        print("TIDE BRONZE -> SILVER COMPLETED")
        print("=" * 100)

    finally:
        spark.stop()


if __name__ == "__main__":
    main()