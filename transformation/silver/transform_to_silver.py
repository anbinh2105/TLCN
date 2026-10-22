# ============================================
# IMPORTS
# ============================================
import csv
import os
import sys
from pathlib import Path
from datetime import datetime, time

from openpyxl import load_workbook

from pyspark.sql import SparkSession
from pyspark.sql import functions as F


# ============================================
# PROJECT CONFIG
# ============================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]


# ============================================
# SPARK
# ============================================

def create_spark():

    python_executable = sys.executable

    # ========================================================
    # PYSPARK PYTHON
    # ========================================================

    # Đảm bảo Spark Driver và Worker
    # sử dụng đúng Python trong virtual environment.
    os.environ["PYSPARK_PYTHON"] = python_executable
    os.environ["PYSPARK_DRIVER_PYTHON"] = python_executable

    # ========================================================
    # HADOOP ON WINDOWS
    # ========================================================

    hadoop_home = r"C:\hadoop"
    hadoop_bin = os.path.join(
        hadoop_home,
        "bin"
    )

    winutils_path = os.path.join(
        hadoop_bin,
        "winutils.exe"
    )

    if not os.path.exists(winutils_path):
        raise FileNotFoundError(
            f"winutils.exe not found: {winutils_path}"
        )

    # Python process
    os.environ["HADOOP_HOME"] = hadoop_home

    # Add C:\hadoop\bin vào PATH nếu chưa có.
    current_path = os.environ.get(
        "PATH",
        ""
    )

    if hadoop_bin.lower() not in current_path.lower():

        os.environ["PATH"] = (
            hadoop_bin
            + os.pathsep
            + current_path
        )

    # ========================================================
    # CREATE SPARK
    # ========================================================

    spark = (
        SparkSession.builder
        .master("local[2]")
        .appName(
            "FloodLakehouseSilverTransformation"
        )
        .config(
            "spark.sql.session.timeZone",
            "Asia/Ho_Chi_Minh"
        )
        .config(
            "spark.pyspark.python",
            python_executable
        )
        .config(
            "spark.pyspark.driver.python",
            python_executable
        )
        .config(
            "spark.hadoop.hadoop.home.dir",
            hadoop_home
        )
        .getOrCreate()
    )

    spark.sparkContext.setLogLevel(
        "WARN"
    )

    print(
        "Spark Python:",
        spark.sparkContext.pythonExec
    )

    print(
        "HADOOP_HOME:",
        os.environ.get("HADOOP_HOME")
    )

    print(
        "winutils.exe:",
        winutils_path
    )

    return spark

# ============================================
# WEATHER
# ============================================

def transform_weather(spark, input_path):

    df = (
        spark.read
        .option("header", True)
        .option("inferSchema", False)
        .csv(str(input_path))
    )

    df = (
        df
        .withColumn(
            "timestamp",
            F.to_timestamp("timestamp")
        )
        .withColumn(
            "temperature_c",
            F.col("temperature").cast("double")
        )
        .withColumn(
            "humidity_pct",
            F.col("humidity").cast("double")
        )
        .withColumn(
            "pressure_hpa",
            F.col("pressure").cast("double")
        )
        .withColumn(
            "wind_speed",
            F.col("wind_speed").cast("double")
        )
        .withColumn(
            "wind_direction",
            F.col("wind_direction").cast("double")
        )
        .withColumn(
            "latitude",
            F.col("latitude").cast("double")
        )
        .withColumn(
            "longitude",
            F.col("longitude").cast("double")
        )
    )

    df = df.withColumn(
        "quality_flag",

        F.when(
            F.col("timestamp").isNull(),
            F.lit("INVALID")
        )

        .when(
            F.col("temperature_c").isNull(),
            F.lit("MISSING")
        )

        .otherwise(
            F.lit("OBSERVED")
        )
    )

    return df.select(
        "timestamp",
        "temperature_c",
        "humidity_pct",
        "pressure_hpa",
        "wind_speed",
        "wind_direction",
        "latitude",
        "longitude",
        "quality_flag"
    )


# ============================================
# RAINFALL
# ============================================

def transform_rainfall(spark, input_path):

    df = (
        spark.read
        .option("header", True)
        .option("inferSchema", False)
        .csv(str(input_path))
    )

    df = (
        df
        .withColumn(
            "timestamp",
            F.to_timestamp("timestamp")
        )
        .withColumn(
            "rainfall_mm",
            F.col("rainfall_mm").cast("double")
        )
        .withColumn(
            "latitude",
            F.col("latitude").cast("double")
        )
        .withColumn(
            "longitude",
            F.col("longitude").cast("double")
        )
        .withColumn(
            "quality_flag",

            F.when(
                F.col("timestamp").isNull(),
                F.lit("INVALID")
            )

            .when(
                F.col("rainfall_mm").isNull(),
                F.lit("MISSING")
            )

            .otherwise(
                F.lit("OBSERVED")
            )
        )
    )

    return df.select(
        "timestamp",
        "rainfall_mm",
        "latitude",
        "longitude",
        "quality_flag"
    )


# ============================================
# RIVER DISCHARGE
# ============================================

def transform_discharge(spark, input_path):

    df = (
        spark.read
        .option("header", True)
        .option("inferSchema", False)
        .csv(str(input_path))
    )

    df = (
        df
        .withColumn(
            "timestamp",
            F.to_timestamp("timestamp")
        )
        .withColumn(
            "discharge_m3s",
            F.col("discharge_m3s").cast("double")
        )
        .withColumn(
            "quality_flag",

            F.when(
                F.col("timestamp").isNull(),
                F.lit("INVALID")
            )

            .when(
                F.col("discharge_m3s").isNull(),
                F.lit("MISSING")
            )

            .otherwise(
                F.lit("OBSERVED")
            )
        )
    )

    return df.select(
        "timestamp",
        "station",
        "discharge_m3s",
        "quality_flag"
    )


# ============================================
# TIDE - XLSX READER
# ============================================

def read_tide_xlsx(file_path):

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    ws = workbook.active

    rows = []

    for row in range(2, ws.max_row + 1):

        date_cell = ws.cell(row, 1)
        time_cell = ws.cell(row, 2)
        tide_cell = ws.cell(row, 3)

        date_value = date_cell.value
        time_value = time_cell.value
        raw_tide = tide_cell.value

        tide_value = None

        quality_flag = "OBSERVED"
        recovery_method = "NONE"

        # ----------------------------
        # Normal numeric
        # ----------------------------

        if isinstance(raw_tide, (int, float)):

            tide_value = float(raw_tide)

        # ----------------------------
        # Numeric stored as string
        # ----------------------------

        elif isinstance(raw_tide, str):

            try:

                tide_value = float(raw_tide)

            except ValueError:

                quality_flag = "INVALID"

        # ----------------------------
        # Excel date representation
        # Example:
        # 2026-07-02 + d.m -> 2.7
        # ----------------------------

        elif (
            isinstance(raw_tide, datetime)
            and tide_cell.number_format == r"d\.m"
        ):

            day = raw_tide.day
            month = raw_tide.month

            if 1 <= month <= 9:

                tide_value = (
                    float(day)
                    + float(month) / 10
                )

                quality_flag = "RECOVERED"

                recovery_method = (
                    "EXCEL_DATE_REPRESENTATION"
                )

            else:

                quality_flag = "INVALID"

        # ----------------------------
        # Missing
        # ----------------------------

        elif raw_tide is None:

            quality_flag = "MISSING"

        # ----------------------------
        # Unknown representation
        # ----------------------------

        else:

            quality_flag = "INVALID"

        # ----------------------------
        # Date -> string
        # ----------------------------

        if date_value is None:

            date_string = None

        elif isinstance(date_value, datetime):

            date_string = date_value.strftime(
                "%Y-%m-%d"
            )

        else:

            date_string = str(
                date_value
            ).strip()


        # ----------------------------
        # Time -> string
        # ----------------------------

        if time_value is None:

            time_string = None

        elif isinstance(time_value, time):

            time_string = time_value.strftime(
                "%H:%M:%S"
            )

        elif isinstance(time_value, datetime):

            time_string = time_value.strftime(
                "%H:%M:%S"
            )

        else:

            time_string = str(
                time_value
            ).strip()

        rows.append({

            "date": date_string,

            "time": time_string,

            "tide_m": tide_value,

            "quality_flag": quality_flag,

            "recovery_method": recovery_method,

            "source_file": Path(file_path).name
        })

    workbook.close()

    return rows

# ============================================
# TIDE - CSV DELIMITER DETECTION
# ============================================

def detect_csv_delimiter(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace"
    ) as file:

        sample = file.read(4096)

    try:

        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=",;\t|"
        )

        return dialect.delimiter

    except csv.Error:

        return ","


# ============================================
# TIDE - CSV READER
# ============================================

def read_tide_csv(file_path):

    delimiter = detect_csv_delimiter(file_path)

    rows = []

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:

        reader = csv.DictReader(
            file,
            delimiter=delimiter
        )

        for record in reader:

            date_value = record.get("Ngày")
            time_value = record.get("Giờ")
            raw_tide = record.get("Độ cao (m)")

            tide_value = None
            quality_flag = "OBSERVED"
            recovery_method = "NONE"

            # ----------------------------
            # Missing tide
            # ----------------------------

            if (
                raw_tide is None
                or str(raw_tide).strip() == ""
            ):

                quality_flag = "MISSING"

            # ----------------------------
            # Numeric tide
            # ----------------------------

            else:

                try:

                    tide_value = float(
                        str(raw_tide)
                        .strip()
                        .replace(",", ".")
                    )

                except ValueError:

                    quality_flag = "INVALID"

            rows.append({

                "date": (
                    str(date_value).strip()
                    if date_value is not None
                    else None
                ),

                "time": (
                    str(time_value).strip()
                    if time_value is not None
                    else None
                ),

                "tide_m": tide_value,

                "quality_flag": quality_flag,

                "recovery_method": recovery_method,

                "source_file": Path(file_path).name
            })

    return rows

# ============================================
# TIDE TRANSFORMATION
# ============================================

def transform_tide(spark, input_path):

    input_path = Path(input_path)

    extension = input_path.suffix.lower()

    # ----------------------------
    # Read source
    # ----------------------------

    if extension == ".xlsx":

        rows = read_tide_xlsx(
            input_path
        )

    elif extension == ".csv":

        rows = read_tide_csv(
            input_path
        )

    else:

        raise ValueError(
            f"Unsupported Tide file: {input_path}"
        )

    if not rows:

        raise ValueError(
            f"Tide file contains no records: {input_path}"
        )

    # ----------------------------
    # Python rows -> Spark
    # ----------------------------

    df = spark.createDataFrame(rows)

    # ----------------------------
    # Remove completely blank rows
    # ----------------------------
    # Một số XLSX 2021 có max_row lớn hơn vùng dữ liệu thực tế.
    # Nếu date, time và tide_m đều NULL thì đây không phải
    # observation bị thiếu mà là blank worksheet row.

    df = df.filter(
        ~(
            F.col("date").isNull()
            & F.col("time").isNull()
            & F.col("tide_m").isNull()
        )
    )

    # ----------------------------
    # Normalize time
    # ----------------------------
    # CSV 2024-2025 có giờ dạng:
    # 0:00, 1:00, ..., 9:00
    #
    # Chuẩn hóa thành:
    # 00:00, 01:00, ..., 09:00
    #
    # Các giá trị 10:00 -> 23:00 giữ nguyên.

    normalized_time = (
        F.when(
            F.col("time").isNull()
            | (F.trim(F.col("time")) == ""),
            F.lit(None)
        )
        .when(
            F.trim(F.col("time")).rlike(
                r"^\d:\d{2}$"
            ),
            F.concat(
                F.lit("0"),
                F.trim(F.col("time"))
            )
        )
        .otherwise(
            F.trim(F.col("time"))
        )
    )

    df = df.withColumn(
        "time_normalized",
        normalized_time
    )


    # ----------------------------
    # Create datetime string
    # ----------------------------

    datetime_string = (
        F.when(
            F.col("date").isNull()
            | F.col("time_normalized").isNull()
            | (F.trim(F.col("date")) == ""),
            F.lit(None)
        )
        .otherwise(
            F.concat_ws(
                " ",
                F.trim(F.col("date")),
                F.col("time_normalized")
            )
        )
    )

    df = df.withColumn(
        "datetime_raw",
        datetime_string
    )

    # ----------------------------
    # Parse heterogeneous timestamps
    # ----------------------------
    #
    # Tide có nhiều format:
    #
    # yyyy-MM-dd HH:mm:ss
    # yyyy-MM-dd HH:mm
    # yyyy-MM-dd H
    #
    # yyyy/MM/dd HH:mm:ss
    # yyyy/MM/dd HH:mm
    # yyyy/MM/dd H
    #
    # dd/MM/yyyy HH:mm:ss
    # dd/MM/yyyy HH:mm
    #
    # try_to_timestamp:
    # parse lỗi -> NULL, không crash pipeline.

    df = df.withColumn(
        "timestamp",

        F.coalesce(

            # --------------------
            # yyyy-MM-dd
            # --------------------

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy-MM-dd HH:mm:ss")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy-MM-dd HH:mm")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy-MM-dd H")
            ),

            # --------------------
            # yyyy/MM/dd
            # --------------------

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy/MM/dd HH:mm:ss")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy/MM/dd HH:mm")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("yyyy/MM/dd H")
            ),

            # --------------------
            # dd/MM/yyyy
            # --------------------

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("dd/MM/yyyy HH:mm:ss")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("dd/MM/yyyy HH:mm")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("d/M/yyyy HH:mm:ss")
            ),

            F.try_to_timestamp(
                F.col("datetime_raw"),
                F.lit("d/M/yyyy HH:mm")
            )
        )
    )

    # ----------------------------
    # Standardize tide datatype
    # ----------------------------

    df = df.withColumn(
        "tide_m",
        F.col("tide_m").cast("double")
    )

    # ----------------------------
    # Quality flag
    # ----------------------------

    df = df.withColumn(
        "quality_flag",

        F.when(
            F.col("timestamp").isNull(),
            F.lit("INVALID")
        )

        .when(
            F.col("tide_m").isNull(),
            F.lit("MISSING")
        )

        .otherwise(
            F.col("quality_flag")
        )
    )

    # ----------------------------
    # Final Silver schema
    # ----------------------------

    return df.select(
        "timestamp",
        "tide_m",
        "quality_flag",
        "recovery_method",
        "source_file"
    )

# ============================================
# WATER LEVEL - XLSX READER
# ============================================

def read_water_level_xlsx(file_path):

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    ws = workbook.active

    # ----------------------------
    # Read header
    # ----------------------------

    headers = [
        cell.value
        for cell in ws[1]
    ]

    rows = []

    # ----------------------------
    # Read records
    # ----------------------------

    for values in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        record = dict(
            zip(headers, values)
        )

        timestamp_value = record.get(
            "Timestamp (UTC+07:00)"
        )

        raw_value = record.get(
            "Value"
        )

        station_code = record.get(
            "Station Code"
        )

        unit = record.get(
            "Unit"
        )

        # ----------------------------
        # Water level value
        # ----------------------------

        water_level = None

        quality_flag = "OBSERVED"

        if raw_value is None:

            quality_flag = "MISSING"

        else:

            try:

                water_level = float(
                    raw_value
                )

            except (
                ValueError,
                TypeError
            ):

                quality_flag = "INVALID"

        # ----------------------------
        # Timestamp
        # ----------------------------

        if isinstance(
            timestamp_value,
            datetime
        ):

            timestamp_string = (
                timestamp_value.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )

        elif timestamp_value is None:

            timestamp_string = None

            quality_flag = "INVALID"

        else:

            timestamp_string = str(
                timestamp_value
            ).strip()

        rows.append({

            "timestamp_raw":
                timestamp_string,

            "station_code":
                str(station_code)
                if station_code is not None
                else None,

            "water_level_m":
                water_level,

            "unit":
                str(unit)
                if unit is not None
                else None,

            "quality_flag":
                quality_flag,

            "source_file":
                Path(file_path).name
        })

    workbook.close()

    return rows


# ============================================
# WATER LEVEL TRANSFORMATION
# ============================================

def transform_water_level(
    spark,
    input_path
):

    rows = read_water_level_xlsx(
        input_path
    )

    if not rows:

        raise ValueError(
            f"Water Level file contains no records: "
            f"{input_path}"
        )

    df = spark.createDataFrame(
        rows
    )

    # ----------------------------
    # Standardize timestamp
    # ----------------------------

    df = df.withColumn(
        "timestamp",

        F.to_timestamp(
            F.col("timestamp_raw")
        )
    )

    # ----------------------------
    # Standardize datatype
    # ----------------------------

    df = df.withColumn(
        "water_level_m",
        F.col("water_level_m")
        .cast("double")
    )

    # ----------------------------
    # Invalid timestamp
    # ----------------------------

    df = df.withColumn(
        "quality_flag",

        F.when(
            F.col("timestamp").isNull(),
            F.lit("INVALID")
        )

        .otherwise(
            F.col("quality_flag")
        )
    )

    return df.select(
        "timestamp",
        "station_code",
        "water_level_m",
        "unit",
        "quality_flag",
        "source_file"
    )

# ============================================
# HELPER - SHOW RESULT
# ============================================

def show_result(source_name, df):

    print("\n")
    print("=" * 80)
    print(f"SOURCE: {source_name.upper()}")
    print("=" * 80)

    print("\nSCHEMA:")
    df.printSchema()

    print("\nSAMPLE DATA:")
    df.show(
        10,
        truncate=False
    )

    print("\nQUALITY FLAG:")

    (
        df
        .groupBy("quality_flag")
        .count()
        .orderBy("quality_flag")
        .show(
            truncate=False
        )
    )

    total_records = df.count()

    print(
        f"TOTAL RECORDS: {total_records}"
    )

    print("=" * 80)


# ============================================
# MAIN
# ============================================

# def main():

#     print("=" * 80)
#     print("FLOOD LAKEHOUSE")
#     print("BRONZE -> SILVER TRANSFORMATION")
#     print("=" * 80)

#     # ----------------------------------------
#     # Create Spark
#     # ----------------------------------------

#     spark = create_spark()

#     print("\nSpark started successfully.")
#     print(
#         "Spark version:",
#         spark.version
#     )

#     # ----------------------------------------
#     # Input paths
#     # ----------------------------------------

#     weather_path = (
#         PROJECT_ROOT
#         / "data"
#         / "raw"
#         / "weather"
#         / "weather_can_tho_2018_2026.csv"
#     )

#     rainfall_path = (
#         PROJECT_ROOT
#         / "data"
#         / "raw"
#         / "rainfall"
#         / "rainfall_can_tho_2018_2026.csv"
#     )

#     discharge_path = (
#         PROJECT_ROOT
#         / "data"
#         / "raw"
#         / "river_discharge"
#         / "river_discharge_can_tho_2018_2026.csv"
#     )

#     water_level_path = (
#         PROJECT_ROOT
#         / "data"
#         / "raw"
#         / "water_level"
#         / "water_level_can_tho_2018_2026.xlsx"
#     )

#     tide_directory = (
#         PROJECT_ROOT
#         / "data"
#         / "raw"
#         / "tide"
#     )

#     # ----------------------------------------
#     # Output directory
#     # ----------------------------------------

#     output_root = (
#         PROJECT_ROOT
#         / "data"
#         / "temp"
#         / "silver-test"
#     )

#     output_root.mkdir(
#         parents=True,
#         exist_ok=True
#     )

#     # ----------------------------------------
#     # Check paths
#     # ----------------------------------------

#     required_files = {

#         "weather":
#             weather_path,

#         "rainfall":
#             rainfall_path,

#         "river_discharge":
#             discharge_path,

#         "water_level":
#             water_level_path
#     }

#     print("\nChecking input files...")

#     for source, path in required_files.items():

#         if not path.exists():

#             raise FileNotFoundError(
#                 f"{source}: File not found -> {path}"
#             )

#         print(
#             f"[OK] {source}: {path.name}"
#         )

#     if not tide_directory.exists():

#         raise FileNotFoundError(
#             f"Tide directory not found: "
#             f"{tide_directory}"
#         )

#     print(
#         f"[OK] tide directory: "
#         f"{tide_directory}"
#     )

#     # ========================================
#     # WEATHER
#     # ========================================

#     print("\nTransforming WEATHER...")

#     weather_df = transform_weather(
#         spark,
#         weather_path
#     )

#     show_result(
#         "weather",
#         weather_df
#     )

#     weather_output = (
#         output_root
#         / "weather"
#     )

#     (
#         weather_df
#         .write
#         .mode("overwrite")
#         .parquet(
#             str(weather_output)
#         )
#     )

#     print(
#         f"[SAVED] {weather_output}"
#     )

#     # ========================================
#     # RAINFALL
#     # ========================================

#     print("\nTransforming RAINFALL...")

#     rainfall_df = transform_rainfall(
#         spark,
#         rainfall_path
#     )

#     show_result(
#         "rainfall",
#         rainfall_df
#     )

#     rainfall_output = (
#         output_root
#         / "rainfall"
#     )

#     (
#         rainfall_df
#         .write
#         .mode("overwrite")
#         .parquet(
#             str(rainfall_output)
#         )
#     )

#     print(
#         f"[SAVED] {rainfall_output}"
#     )

#     # ========================================
#     # RIVER DISCHARGE
#     # ========================================

#     print(
#         "\nTransforming RIVER DISCHARGE..."
#     )

#     discharge_df = transform_discharge(
#         spark,
#         discharge_path
#     )

#     show_result(
#         "river_discharge",
#         discharge_df
#     )

#     discharge_output = (
#         output_root
#         / "river_discharge"
#     )

#     (
#         discharge_df
#         .write
#         .mode("overwrite")
#         .parquet(
#             str(discharge_output)
#         )
#     )

#     print(
#         f"[SAVED] {discharge_output}"
#     )

#     # ========================================
#     # WATER LEVEL
#     # ========================================

#     print(
#         "\nTransforming WATER LEVEL..."
#     )

#     water_level_df = transform_water_level(
#         spark,
#         water_level_path
#     )

#     show_result(
#         "water_level",
#         water_level_df
#     )

#     water_level_output = (
#         output_root
#         / "water_level"
#     )

#     (
#         water_level_df
#         .write
#         .mode("overwrite")
#         .parquet(
#             str(water_level_output)
#         )
#     )

#     print(
#         f"[SAVED] {water_level_output}"
#     )

#     # ========================================
#     # TIDE
#     # ========================================

#     print("\nTransforming TIDE...")

#     tide_files = sorted([
#         path
#         for path in tide_directory.iterdir()
#         if path.suffix.lower()
#         in {".xlsx", ".csv"}
#     ])

#     if not tide_files:

#         raise FileNotFoundError(
#             "No Tide CSV/XLSX files found."
#         )

#     print(
#         f"Tide files found: "
#         f"{len(tide_files)}"
#     )

#     tide_dataframes = []

#     for tide_file in tide_files:

#         print(
#             f"  Processing: "
#             f"{tide_file.name}"
#         )

#         try:

#             tide_df = transform_tide(
#                 spark,
#                 tide_file
#             )

#             tide_dataframes.append(
#                 tide_df
#             )

#         except Exception as error:

#             print(
#                 f"  [ERROR] "
#                 f"{tide_file.name}: "
#                 f"{error}"
#             )

#             raise

#     # ----------------------------------------
#     # Union all Tide files
#     # ----------------------------------------

#     tide_all_df = tide_dataframes[0]

#     for tide_df in tide_dataframes[1:]:

#         tide_all_df = (
#             tide_all_df
#             .unionByName(
#                 tide_df,
#                 allowMissingColumns=True
#             )
#         )

#     show_result(
#         "tide",
#         tide_all_df
#     )

#     # ----------------------------------------
#     # Tide recovery statistics
#     # ----------------------------------------

#     print(
#         "\nTIDE RECOVERY METHOD:"
#     )

#     (
#         tide_all_df
#         .groupBy(
#             "recovery_method"
#         )
#         .count()
#         .orderBy(
#             "recovery_method"
#         )
#         .show(
#             truncate=False
#         )
#     )

#     tide_output = (
#         output_root
#         / "tide"
#     )

#     (
#         tide_all_df
#         .write
#         .mode("overwrite")
#         .parquet(
#             str(tide_output)
#         )
#     )

#     print(
#         f"[SAVED] {tide_output}"
#     )

#     # ========================================
#     # FINAL SUMMARY
#     # ========================================

#     print("\n")
#     print("=" * 80)
#     print("SILVER TRANSFORMATION COMPLETED")
#     print("=" * 80)

#     print(
#         f"Weather        : "
#         f"{weather_df.count()} records"
#     )

#     print(
#         f"Rainfall       : "
#         f"{rainfall_df.count()} records"
#     )

#     print(
#         f"River Discharge: "
#         f"{discharge_df.count()} records"
#     )

#     print(
#         f"Water Level    : "
#         f"{water_level_df.count()} records"
#     )

#     print(
#         f"Tide           : "
#         f"{tide_all_df.count()} records"
#     )

#     print(
#         "\nSilver test output:"
#     )

#     print(output_root)

#     print("=" * 80)

#     # ----------------------------------------
#     # Stop Spark
#     # ----------------------------------------

#     spark.stop()


# # ============================================
# # ENTRY POINT
# # ============================================

# if __name__ == "__main__":
#     main()

# ============================================
# Kiểm tra riêng tide
# ============================================
def main_tide_only():

    print("=" * 80)
    print("TIDE ONLY - BRONZE -> SILVER")
    print("=" * 80)

    spark = create_spark()

    tide_directory = (
        PROJECT_ROOT
        / "data"
        / "raw"
        / "tide"
    )

    output_path = (
        PROJECT_ROOT
        / "data"
        / "temp"
        / "silver-test"
        / "tide"
    )

    tide_files = sorted([
        path
        for path in tide_directory.iterdir()
        if path.suffix.lower() in {".xlsx", ".csv"}
    ])

    print(f"Tide files found: {len(tide_files)}")

    tide_dataframes = []

    for tide_file in tide_files:

        print(f"Processing: {tide_file.name}")

        tide_df = transform_tide(
            spark,
            tide_file
        )

        tide_dataframes.append(
            tide_df
        )

    tide_all_df = tide_dataframes[0]

    for tide_df in tide_dataframes[1:]:

        tide_all_df = tide_all_df.unionByName(
            tide_df,
            allowMissingColumns=True
        )

    # ----------------------------
    # Result
    # ----------------------------

    show_result(
        "tide",
        tide_all_df
    )

    print("\nTIDE RECOVERY METHOD:")

    (
        tide_all_df
        .groupBy("recovery_method")
        .count()
        .orderBy("recovery_method")
        .show(truncate=False)
    )

    # ----------------------------
    # INVALID BY FILE
    # ----------------------------

    print("\nINVALID BY SOURCE FILE:")

    (
        tide_all_df
        .filter(
            F.col("quality_flag") == "INVALID"
        )
        .groupBy("source_file")
        .count()
        .orderBy(
            F.desc("count")
        )
        .show(
            100,
            truncate=False
        )
    )

    # ----------------------------
    # Write Silver test
    # ----------------------------

    (
        tide_all_df
        .write
        .mode("overwrite")
        .parquet(
            str(output_path)
        )
    )

    print(f"\n[SAVED] {output_path}")

    spark.stop()


if __name__ == "__main__":
    main_tide_only()