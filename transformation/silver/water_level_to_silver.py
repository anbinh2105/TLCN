"""
Bronze -> Silver transformation for Can Tho Water Level.

Architecture:
    MinIO Bronze XLSX
        -> OpenPyXL adapter
        -> Spark DataFrame
        -> Standardization / Quality
        -> Apache Iceberg Silver

Silver table:
    lakehouse.silver.water_level

Important:
    - Source timestamps are UTC+07:00.
    - Native temporal resolution is 15 minutes.
    - Silver preserves native resolution.
    - No interpolation, resampling, or imputation.
"""

import io
import os
from typing import List, Tuple

from dotenv import load_dotenv
from minio import Minio
from openpyxl import load_workbook

from pyspark.sql import DataFrame
from pyspark.sql import functions as F
from pyspark.sql.types import (
    DoubleType,
    StringType,
    StructField,
    StructType,
    TimestampType,
)

from transformation.silver.common import (
    QUALITY_OBSERVED,
    QUALITY_MISSING,
    QUALITY_INVALID,
    create_spark_session,
    get_silver_table,
    ensure_silver_namespace,
    add_silver_metadata,
    remove_exact_duplicates,
    print_dataframe_summary,
    stop_spark,
)


# ============================================================
# Configuration
# ============================================================

load_dotenv()

SOURCE = "water_level"

MINIO_ENDPOINT = os.getenv(
    "MINIO_ENDPOINT",
    "localhost:9000",
)

MINIO_ACCESS_KEY = os.getenv(
    "MINIO_ACCESS_KEY",
    "admin",
)

MINIO_SECRET_KEY = os.getenv(
    "MINIO_SECRET_KEY",
    "admin12345",
)

MINIO_BUCKET = os.getenv(
    "MINIO_BUCKET",
    "flood-lakehouse",
)

MINIO_SECURE = (
    os.getenv("MINIO_SECURE", "false")
    .strip()
    .lower()
    == "true"
)

BRONZE_PREFIX = "bronze/water_level/"

DATA_SHEET = "Data"


# ============================================================
# Canonical raw adapter schema
# ============================================================

RAW_SCHEMA = StructType(
    [
        StructField("country_code", StringType(), True),
        StructField("station_code", StringType(), True),
        StructField("parameter", StringType(), True),
        StructField("label", StringType(), True),
        StructField("timestamp", TimestampType(), True),
        StructField("water_level_m", DoubleType(), True),
        StructField("unit", StringType(), True),
        StructField("observation_type", StringType(), True),
        StructField("approval_level", StringType(), True),
        StructField("grade", StringType(), True),
        StructField("_bronze_source_file", StringType(), False),
    ]
)


# ============================================================
# MinIO client
# ============================================================

def create_minio_client() -> Minio:
    return Minio(
        MINIO_ENDPOINT,
        access_key=MINIO_ACCESS_KEY,
        secret_key=MINIO_SECRET_KEY,
        secure=MINIO_SECURE,
    )


# ============================================================
# Find Bronze XLSX objects
# ============================================================

def list_water_level_objects(
    client: Minio,
) -> List[str]:

    objects = client.list_objects(
        MINIO_BUCKET,
        prefix=BRONZE_PREFIX,
        recursive=True,
    )

    xlsx_objects = [
        obj.object_name
        for obj in objects
        if obj.object_name.lower().endswith(".xlsx")
    ]

    xlsx_objects.sort()

    if not xlsx_objects:
        raise RuntimeError(
            "No Water Level XLSX files found "
            f"under s3://{MINIO_BUCKET}/{BRONZE_PREFIX}"
        )

    return xlsx_objects


# ============================================================
# Safe value conversion
# ============================================================

def clean_string(value):
    """
    Preserve source text while normalizing Python values.
    """

    if value is None:
        return None

    value = str(value).strip()

    return value if value else None


def safe_float(value):
    """
    Convert source value to float.

    Invalid values become None and will later receive
    a quality flag instead of causing the pipeline to fail.
    """

    if value is None:
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# ============================================================
# Read one XLSX object
# ============================================================

def read_water_level_xlsx(
    client: Minio,
    object_name: str,
) -> List[Tuple]:

    print()
    print(f"[INFO] Reading XLSX: {object_name}")

    response = client.get_object(
        MINIO_BUCKET,
        object_name,
    )

    try:
        file_bytes = response.read()
    finally:
        response.close()
        response.release_conn()

    workbook = load_workbook(
        io.BytesIO(file_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        if DATA_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{DATA_SHEET}' not found in "
                f"{object_name}. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[DATA_SHEET]

        rows = worksheet.iter_rows(
            min_row=2,
            values_only=True,
        )

        records = []

        for row in rows:

            # Expected source schema has 10 columns.
            if len(row) < 10:
                continue

            (
                country_code,
                station_code,
                parameter,
                label,
                timestamp,
                value,
                unit,
                observed_or_estimated,
                approval_level,
                grade,
            ) = row[:10]

            # Ignore completely empty spreadsheet rows.
            if all(v is None for v in row[:10]):
                continue

            records.append(
                (
                    clean_string(country_code),
                    clean_string(station_code),
                    clean_string(parameter),
                    clean_string(label),
                    timestamp,
                    safe_float(value),
                    clean_string(unit),
                    clean_string(observed_or_estimated),
                    clean_string(approval_level),
                    clean_string(grade),
                    f"s3a://{MINIO_BUCKET}/{object_name}",
                )
            )

        print(
            f"[INFO] Rows extracted: {len(records)}"
        )

        return records

    finally:
        workbook.close()


# ============================================================
# Bronze XLSX -> Spark DataFrame
# ============================================================

def read_water_level_bronze(
    spark,
) -> DataFrame:

    print()
    print("[INFO] Reading Bronze Water Level")

    client = create_minio_client()

    object_names = list_water_level_objects(
        client
    )

    print(
        f"[INFO] XLSX objects found: "
        f"{len(object_names)}"
    )

    all_records = []

    for object_name in object_names:
        records = read_water_level_xlsx(
            client,
            object_name,
        )

        all_records.extend(records)

    if not all_records:
        raise RuntimeError(
            "No Water Level records were extracted."
        )

    df = spark.createDataFrame(
        all_records,
        schema=RAW_SCHEMA,
    )

    return df


# ============================================================
# Validate source schema
# ============================================================

def validate_required_columns(
    df: DataFrame,
) -> None:

    required_columns = [
        "station_code",
        "parameter",
        "timestamp",
        "water_level_m",
        "unit",
    ]

    missing = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing:
        raise ValueError(
            "Water Level is missing required "
            f"columns: {missing}"
        )

    print(
        "[PASS] Water Level required columns exist"
    )


# ============================================================
# Standardization
# ============================================================

def standardize_water_level(
    df: DataFrame,
) -> DataFrame:
    """
    Standardize text and numeric representation.

    Timestamp already arrives from OpenPyXL as Python
    datetime and Spark stores it as TimestampType.

    Source timestamps represent UTC+07:00 local time.
    Spark session timezone is Asia/Ho_Chi_Minh.
    """

    return (
        df

        .withColumn(
            "country_code",
            F.trim(F.col("country_code")),
        )

        .withColumn(
            "station_code",
            F.trim(F.col("station_code")),
        )

        .withColumn(
            "parameter",
            F.trim(F.col("parameter")),
        )

        .withColumn(
            "label",
            F.trim(F.col("label")),
        )

        .withColumn(
            "water_level_m",
            F.col("water_level_m").cast(
                DoubleType()
            ),
        )

        .withColumn(
            "unit",
            F.lower(
                F.trim(F.col("unit"))
            ),
        )

        .withColumn(
            "observation_type",
            F.trim(
                F.col("observation_type")
            ),
        )

        .withColumn(
            "approval_level",
            F.trim(
                F.col("approval_level")
            ),
        )

        .withColumn(
            "grade",
            F.trim(F.col("grade")),
        )
    )


# ============================================================
# Quality flag
# ============================================================

def add_water_level_quality_flag(
    df: DataFrame,
) -> DataFrame:
    """
    Row-level quality classification.

    INVALID:
        timestamp missing/invalid
        parameter is not Water Level
        unit is present but not metre

    MISSING:
        station code missing
        water level missing

    OBSERVED:
        structurally valid source observation

    NOTE:
        Approval Level and Grade are preserved as
        source metadata. We do NOT automatically mark
        "Raw - Not Yet Reviewed" or "Unverified data"
        as invalid because those are source review
        statuses, not proof that the measurement itself
        is unusable.
    """

    station_missing = (
        F.col("station_code").isNull()
        | (F.trim(F.col("station_code")) == "")
    )

    parameter_invalid = (
        F.col("parameter").isNotNull()
        & (
            F.lower(F.trim(F.col("parameter")))
            != F.lit("water level")
        )
    )

    unit_invalid = (
        F.col("unit").isNotNull()
        & (
            F.lower(F.trim(F.col("unit")))
            != F.lit("m")
        )
    )

    return df.withColumn(
        "quality_flag",

        F.when(
            F.col("timestamp").isNull(),
            F.lit(QUALITY_INVALID),
        )

        .when(
            parameter_invalid | unit_invalid,
            F.lit(QUALITY_INVALID),
        )

        .when(
            station_missing
            | F.col("water_level_m").isNull(),
            F.lit(QUALITY_MISSING),
        )

        .otherwise(
            F.lit(QUALITY_OBSERVED),
        ),
    )


# ============================================================
# Canonical Silver columns
# ============================================================

def select_water_level_columns(
    df: DataFrame,
) -> DataFrame:

    return df.select(
        "timestamp",
        "station_code",
        "water_level_m",

        # Source measurement metadata
        "country_code",
        "parameter",
        "label",
        "unit",
        "observation_type",
        "approval_level",
        "grade",

        # Lakehouse metadata
        "quality_flag",
        "source",
        "source_file",
        "silver_processed_at",
    )


# ============================================================
# Complete transformation
# ============================================================

def transform_water_level(
    df: DataFrame,
) -> DataFrame:

    validate_required_columns(df)

    df = standardize_water_level(df)

    # Exact duplicates only.
    # Do not arbitrarily resolve conflicting
    # station/timestamp observations.
    df = remove_exact_duplicates(df)

    df = add_water_level_quality_flag(df)

    df = add_silver_metadata(
        df,
        source=SOURCE,
        source_file_column="_bronze_source_file",
    )

    return select_water_level_columns(df)


# ============================================================
# Write Iceberg
# ============================================================

def write_water_level_silver(
    df: DataFrame,
) -> None:

    silver_table = get_silver_table(
        SOURCE
    )

    print()
    print("[INFO] Writing Water Level Silver")
    print(
        f"[INFO] Iceberg table: {silver_table}"
    )

    (
        df.writeTo(silver_table)
        .using("iceberg")
        .tableProperty(
            "format-version",
            "2",
        )
        .createOrReplace()
    )

    print(
        "[PASS] Water Level Silver "
        "Iceberg table written"
    )


# ============================================================
# Validation
# ============================================================

def validate_water_level_silver(
    spark,
    silver_table: str,
) -> None:

    df = spark.table(silver_table)

    print()
    print("=" * 70)
    print("SILVER WATER LEVEL VALIDATION")
    print("=" * 70)

    # --------------------------------------------------------
    # Total
    # --------------------------------------------------------

    total_rows = df.count()

    print(
        f"Total rows: {total_rows}"
    )

    if total_rows == 0:
        raise ValueError(
            "Silver Water Level contains zero rows."
        )

    # --------------------------------------------------------
    # Schema
    # --------------------------------------------------------

    print()
    print("Schema:")
    df.printSchema()

    # --------------------------------------------------------
    # Quality flags
    # --------------------------------------------------------

    print()
    print("Quality flags:")

    (
        df
        .groupBy("quality_flag")
        .count()
        .orderBy("quality_flag")
        .show(truncate=False)
    )

    # --------------------------------------------------------
    # Timestamp range
    # --------------------------------------------------------

    print()
    print("Timestamp range:")

    (
        df
        .select(
            F.min("timestamp").alias(
                "min_timestamp"
            ),
            F.max("timestamp").alias(
                "max_timestamp"
            ),
        )
        .show(truncate=False)
    )

    # --------------------------------------------------------
    # Water level statistics
    # --------------------------------------------------------

    print()
    print("Water level statistics (m):")

    (
        df
        .select(
            F.min("water_level_m").alias(
                "min_water_level_m"
            ),
            F.max("water_level_m").alias(
                "max_water_level_m"
            ),
            F.avg("water_level_m").alias(
                "avg_water_level_m"
            ),
        )
        .show(truncate=False)
    )

    # --------------------------------------------------------
    # Nulls
    # --------------------------------------------------------

    print()
    print("Null counts:")

    check_columns = [
        "timestamp",
        "station_code",
        "water_level_m",
        "unit",
    ]

    expressions = [
        F.sum(
            F.when(
                F.col(column).isNull(),
                1,
            ).otherwise(0)
        ).alias(column)
        for column in check_columns
    ]

    df.select(
        *expressions
    ).show(truncate=False)

    # --------------------------------------------------------
    # Station
    # --------------------------------------------------------

    print()
    print("Stations:")

    (
        df
        .groupBy("station_code")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Units
    # --------------------------------------------------------

    print()
    print("Units:")

    (
        df
        .groupBy("unit")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Parameter
    # --------------------------------------------------------

    print()
    print("Parameters:")

    (
        df
        .groupBy("parameter")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Observation type
    # --------------------------------------------------------

    print()
    print("Observed / Estimated:")

    (
        df
        .groupBy("observation_type")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Approval levels
    # --------------------------------------------------------

    print()
    print("Approval levels:")

    (
        df
        .groupBy("approval_level")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Grades
    # --------------------------------------------------------

    print()
    print("Grades:")

    (
        df
        .groupBy("grade")
        .count()
        .orderBy(F.desc("count"))
        .show(
            20,
            truncate=False,
        )
    )

    # --------------------------------------------------------
    # Duplicate business keys
    # --------------------------------------------------------

    print()
    print(
        "Duplicate station + timestamp:"
    )

    duplicate_df = (
        df
        .filter(
            F.col("timestamp").isNotNull()
        )
        .groupBy(
            "station_code",
            "timestamp",
        )
        .count()
        .filter(
            F.col("count") > 1
        )
    )

    duplicate_groups = (
        duplicate_df.count()
    )

    print(
        "Duplicate station-timestamp groups: "
        f"{duplicate_groups}"
    )

    if duplicate_groups > 0:
        (
            duplicate_df
            .orderBy(
                "station_code",
                "timestamp",
            )
            .show(
                20,
                truncate=False,
            )
        )

    # --------------------------------------------------------
    # 15-minute grain check
    # --------------------------------------------------------

    print()
    print("15-minute alignment:")

    misaligned_count = (
        df
        .filter(
            F.col("timestamp").isNotNull()
        )
        .filter(
            (F.minute("timestamp") % 15 != 0)
            | (F.second("timestamp") != 0)
        )
        .count()
    )

    print(
        "Records not aligned to 15-minute "
        f"intervals: {misaligned_count}"
    )

    # --------------------------------------------------------
    # Source files
    # --------------------------------------------------------

    print()
    print("Source files:")

    (
        df
        .select("source_file")
        .distinct()
        .show(
            20,
            truncate=False,
        )
    )

    print("=" * 70)


# ============================================================
# Main
# ============================================================

def main():

    spark = None

    try:

        print()
        print("=" * 70)
        print("CAN THO FLOOD LAKEHOUSE")
        print(
            "BRONZE -> SILVER : WATER LEVEL"
        )
        print("=" * 70)

        # ----------------------------------------------------
        # Spark
        # ----------------------------------------------------

        spark = create_spark_session(
            app_name="water-level-to-silver",
        )

        ensure_silver_namespace(
            spark
        )

        # ----------------------------------------------------
        # XLSX adapter -> Spark
        # ----------------------------------------------------

        bronze_df = read_water_level_bronze(
            spark
        )

        print_dataframe_summary(
            bronze_df,
            "WATER LEVEL BRONZE ADAPTER",
        )

        # ----------------------------------------------------
        # Transform
        # ----------------------------------------------------

        silver_df = transform_water_level(
            bronze_df
        )

        print_dataframe_summary(
            silver_df,
            "WATER LEVEL SILVER BEFORE WRITE",
        )

        # ----------------------------------------------------
        # Write
        # ----------------------------------------------------

        write_water_level_silver(
            silver_df
        )

        # ----------------------------------------------------
        # Persisted validation
        # ----------------------------------------------------

        silver_table = get_silver_table(
            SOURCE
        )

        validate_water_level_silver(
            spark,
            silver_table,
        )

        print()
        print("=" * 70)
        print(
            "WATER LEVEL BRONZE -> SILVER: PASS"
        )
        print("=" * 70)

    except Exception as exc:

        print()
        print("=" * 70)
        print(
            "WATER LEVEL BRONZE -> SILVER: FAILED"
        )
        print("=" * 70)
        print(
            f"{type(exc).__name__}: {exc}"
        )

        raise

    finally:
        stop_spark(spark)


if __name__ == "__main__":
    main()