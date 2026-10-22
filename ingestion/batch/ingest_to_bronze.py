from pathlib import Path
from datetime import datetime
import hashlib
import json
import os
import csv

from dotenv import load_dotenv
from minio import Minio
from openpyxl import load_workbook


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

RAW_ROOT = PROJECT_ROOT / "data" / "raw"

METADATA_DIR = PROJECT_ROOT / "data" / "metadata"

INGESTION_LOG = (
    METADATA_DIR
    / "ingestion_log.jsonl"
)


SOURCES = [
    "weather",
    "rainfall",
    "water_level",
    "river_discharge",
    "tide",
]


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
}


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv(PROJECT_ROOT / ".env")

MINIO_ENDPOINT = os.getenv(
    "MINIO_ENDPOINT",
    "localhost:9000"
)

MINIO_ACCESS_KEY = os.getenv(
    "MINIO_ACCESS_KEY"
)

MINIO_SECRET_KEY = os.getenv(
    "MINIO_SECRET_KEY"
)

MINIO_BUCKET = os.getenv(
    "MINIO_BUCKET",
    "flood-lakehouse"
)


# ============================================================
# MINIO CLIENT
# ============================================================

def create_minio_client():

    if not MINIO_ACCESS_KEY:
        raise ValueError(
            "MINIO_ACCESS_KEY is missing."
        )

    if not MINIO_SECRET_KEY:
        raise ValueError(
            "MINIO_SECRET_KEY is missing."
        )

    return Minio(
        MINIO_ENDPOINT,
        access_key=MINIO_ACCESS_KEY,
        secret_key=MINIO_SECRET_KEY,
        secure=False,
    )


def ensure_bucket(client):

    if not client.bucket_exists(
        MINIO_BUCKET
    ):
        client.make_bucket(
            MINIO_BUCKET
        )

        print(
            f"Created bucket: "
            f"{MINIO_BUCKET}"
        )

    else:
        print(
            f"Bucket exists: "
            f"{MINIO_BUCKET}"
        )


# ============================================================
# CHECKSUM
# ============================================================

def calculate_sha256(file_path):

    sha256 = hashlib.sha256()

    with open(file_path, "rb") as file:

        while True:

            chunk = file.read(
                1024 * 1024
            )

            if not chunk:
                break

            sha256.update(chunk)

    return sha256.hexdigest()


# ============================================================
# RECORD COUNT
# ============================================================

def count_csv_records(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:

        reader = csv.reader(file)

        row_count = sum(
            1
            for row in reader
            if any(
                str(value).strip()
                for value in row
            )
        )

    # subtract header
    return max(
        row_count - 1,
        0
    )


def count_xlsx_records(file_path):

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    total_rows = 0

    for row in worksheet.iter_rows(
        values_only=True
    ):

        if any(
            value is not None
            for value in row
        ):
            total_rows += 1

    workbook.close()

    # subtract header
    return max(
        total_rows - 1,
        0
    )


def count_records(file_path):

    extension = (
        file_path
        .suffix
        .lower()
    )

    if extension == ".csv":
        return count_csv_records(
            file_path
        )

    if extension == ".xlsx":
        return count_xlsx_records(
            file_path
        )

    return None


# ============================================================
# INGESTION LOG
# ============================================================

def load_successful_checksums():

    checksums = set()

    if not INGESTION_LOG.exists():
        return checksums

    with open(
        INGESTION_LOG,
        "r",
        encoding="utf-8"
    ) as file:

        for line in file:

            line = line.strip()

            if not line:
                continue

            try:
                record = json.loads(
                    line
                )

            except json.JSONDecodeError:
                continue

            if (
                record.get("status")
                == "SUCCESS"
                and record.get("checksum")
            ):
                checksums.add(
                    record["checksum"]
                )

    return checksums


def write_ingestion_log(record):

    METADATA_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        INGESTION_LOG,
        "a",
        encoding="utf-8"
    ) as file:

        file.write(
            json.dumps(
                record,
                ensure_ascii=False
            )
            + "\n"
        )


# ============================================================
# FILE DISCOVERY
# ============================================================

def discover_files(source):

    source_dir = (
        RAW_ROOT
        / source
    )

    if not source_dir.exists():
        return []

    files = [
        file_path
        for file_path
        in source_dir.rglob("*")
        if (
            file_path.is_file()
            and file_path.suffix.lower()
            in SUPPORTED_EXTENSIONS
        )
    ]

    return sorted(files)


# ============================================================
# OBJECT PATH
# ============================================================

def build_object_name(
    source,
    file_name,
    ingestion_date,
    batch_id
):

    return (
        f"bronze/"
        f"{source}/"
        f"ingestion_date={ingestion_date}/"
        f"batch_id={batch_id}/"
        f"{file_name}"
    )


# ============================================================
# INGEST ONE FILE
# ============================================================

def ingest_file(
    client,
    source,
    file_path,
    batch_id,
    ingestion_date,
    successful_checksums
):

    ingestion_timestamp = (
        datetime.now()
        .astimezone()
        .isoformat()
    )

    file_size = (
        file_path.stat().st_size
    )

    checksum = calculate_sha256(
        file_path
    )

    print()
    print(
        f"Source : {source}"
    )

    print(
        f"File   : {file_path.name}"
    )

    print(
        f"SHA256 : {checksum}"
    )

    # --------------------------------------------------------
    # IDEMPOTENCY CHECK
    # --------------------------------------------------------

    if checksum in successful_checksums:

        print(
            "Status : SKIPPED "
            "(checksum already ingested)"
        )

        record = {
            "batch_id": batch_id,
            "source": source,
            "source_file": file_path.name,
            "object_path": None,
            "ingestion_timestamp":
                ingestion_timestamp,
            "file_size": file_size,
            "checksum": checksum,
            "record_count": None,
            "status": "SKIPPED",
        }

        write_ingestion_log(
            record
        )

        return "SKIPPED"

    # --------------------------------------------------------
    # RECORD COUNT
    # --------------------------------------------------------

    try:

        record_count = count_records(
            file_path
        )

    except Exception as error:

        print(
            f"Record count warning: "
            f"{error}"
        )

        record_count = None

    # --------------------------------------------------------
    # OBJECT NAME
    # --------------------------------------------------------

    object_name = build_object_name(
        source=source,
        file_name=file_path.name,
        ingestion_date=ingestion_date,
        batch_id=batch_id,
    )

    # --------------------------------------------------------
    # UPLOAD
    # --------------------------------------------------------

    try:

        client.fput_object(
            MINIO_BUCKET,
            object_name,
            str(file_path),
        )

        status = "SUCCESS"

        successful_checksums.add(
            checksum
        )

        print(
            f"Object : {object_name}"
        )

        print(
            f"Records: {record_count}"
        )

        print(
            "Status : SUCCESS"
        )

    except Exception as error:

        status = "FAILED"

        print(
            f"Status : FAILED"
        )

        print(
            f"Error  : {error}"
        )

    # --------------------------------------------------------
    # METADATA
    # --------------------------------------------------------

    record = {
        "batch_id": batch_id,
        "source": source,
        "source_file": file_path.name,
        "object_path": object_name,
        "ingestion_timestamp":
            ingestion_timestamp,
        "file_size": file_size,
        "checksum": checksum,
        "record_count": record_count,
        "status": status,
    }

    write_ingestion_log(
        record
    )

    return status


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 100)
    print("BATCH INGESTION TO MINIO BRONZE")
    print("=" * 100)

    run_time = (
        datetime.now()
        .astimezone()
    )

    batch_id = run_time.strftime(
        "%Y%m%d_%H%M%S"
    )

    ingestion_date = (
        run_time
        .date()
        .isoformat()
    )

    print(
        f"Batch ID       : {batch_id}"
    )

    print(
        f"Ingestion date : "
        f"{ingestion_date}"
    )

    print(
        f"Bucket         : "
        f"{MINIO_BUCKET}"
    )

    client = create_minio_client()

    ensure_bucket(
        client
    )

    successful_checksums = (
        load_successful_checksums()
    )

    print(
        "Known successful checksums: "
        f"{len(successful_checksums)}"
    )

    summary = {
        "SUCCESS": 0,
        "SKIPPED": 0,
        "FAILED": 0,
    }

    discovered_files = 0

    for source in SOURCES:

        files = discover_files(
            source
        )

        print()
        print("=" * 100)
        print(
            f"SOURCE: {source}"
        )
        print("=" * 100)

        print(
            f"Files discovered: "
            f"{len(files)}"
        )

        discovered_files += len(files)

        for file_path in files:

            status = ingest_file(
                client=client,
                source=source,
                file_path=file_path,
                batch_id=batch_id,
                ingestion_date=ingestion_date,
                successful_checksums=
                    successful_checksums,
            )

            summary[status] += 1

    print()
    print("=" * 100)
    print("INGESTION SUMMARY")
    print("=" * 100)

    print(
        f"Batch ID         : "
        f"{batch_id}"
    )

    print(
        f"Discovered files : "
        f"{discovered_files}"
    )

    print(
        f"SUCCESS          : "
        f"{summary['SUCCESS']}"
    )

    print(
        f"SKIPPED          : "
        f"{summary['SKIPPED']}"
    )

    print(
        f"FAILED           : "
        f"{summary['FAILED']}"
    )

    print()
    print(
        f"Ingestion log: "
        f"{INGESTION_LOG}"
    )


if __name__ == "__main__":
    main()