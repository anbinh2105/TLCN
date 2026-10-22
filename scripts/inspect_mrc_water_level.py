
from pathlib import Path
from openpyxl import load_workbook

DATA_DIR = Path("data/raw/water_level")

KEYWORDS = (
    "datum", "gauge", "zero", "reference", "elevation",
    "station", "water level", "level", "unit", "time zone",
    "timezone", "mốc", "cao độ", "hệ cao độ", "trạm",
)

files = sorted(
    p for p in DATA_DIR.rglob("*")
    if p.suffix.lower() in {".xlsx", ".xlsm"}
    and not p.name.startswith("~$")
)

print(f"Found {len(files)} Excel files")

for path in files:
    print("\n" + "=" * 80)
    print("FILE:", path)

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        print("SHEETS:", wb.sheetnames)

        props = wb.properties
        print("TITLE:", props.title)
        print("SUBJECT:", props.subject)
        print("DESCRIPTION:", props.description)
        print("KEYWORDS:", props.keywords)

        for ws in wb.worksheets:
            print(f"\nSHEET: {ws.title}")
            print("SIZE:", ws.max_row, "rows x", ws.max_column, "columns")

            # Inspect the first 15 rows, up to 12 columns.
            for row in ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 15),
                max_col=min(ws.max_column, 12),
                values_only=True,
            ):
                print(row)

            # Search text cells in a limited sample.
            matches = 0
            for row in ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 100),
                max_col=min(ws.max_column, 20),
                values_only=True,
            ):
                for value in row:
                    if isinstance(value, str):
                        if any(k in value.lower() for k in KEYWORDS):
                            print("METADATA CANDIDATE:", value[:300])
                            matches += 1
                            if matches >= 20:
                                break
                if matches >= 20:
                    break

        wb.close()

    except Exception as exc:
        print("ERROR:", type(exc).__name__, str(exc))
